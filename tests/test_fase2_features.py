"""Tests unitarios para las funcionalidades de la Fase 2:
1. Módulo de alertas proactivas por Telegram Bot (quiebres, seguridad, cierre).
2. Sincronización completa de catálogo SQLite hacia Productos.xlsx (desacople y flush).
"""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from src.orders_store import OrderStore
from src.product_store import ProductStore
from src.stock_manager import StockManager
from src.telegram_notifier import (
    cargar_config_telegram,
    enviar_mensaje_telegram,
    guardar_config_telegram,
    notificar_cierre_jornada_telegram,
    notificar_quiebre_stock_telegram,
    notificar_seguridad_bloqueo_telegram,
)
from src.user_store import UserStore


class TestFase2Features(unittest.TestCase):
    def test_telegram_config_carga_y_guarda_correctamente(self):
        """Verifica la persistencia y carga de la configuración de Telegram."""
        with tempfile.TemporaryDirectory() as tmp:
            tmp_root = Path(tmp)
            # 1. Config por defecto
            cfg_def = cargar_config_telegram(tmp_root)
            self.assertEqual(cfg_def["bot_token"], "")
            self.assertEqual(cfg_def["chat_id"], "")
            self.assertFalse(cfg_def["activo"])
            self.assertTrue(cfg_def["notificar_quiebres"])

            # 2. Guardar nueva configuración
            nueva = {
                "bot_token": "123456:TEST_TOKEN",
                "chat_id": "-100987654321",
                "activo": True,
                "notificar_quiebres": True,
                "notificar_seguridad": True,
                "notificar_cierre": True,
            }
            guardar_config_telegram(nueva, tmp_root)

            # 3. Recargar
            cfg_cargada = cargar_config_telegram(tmp_root)
            self.assertEqual(cfg_cargada["bot_token"], "123456:TEST_TOKEN")
            self.assertEqual(cfg_cargada["chat_id"], "-100987654321")
            self.assertTrue(cfg_cargada["activo"])

    def test_enviar_mensaje_telegram_valida_credenciales(self):
        """Valida que enviar mensaje sin credenciales falle controladamente."""
        exito, detalle = enviar_mensaje_telegram("Hola", config={"bot_token": "", "chat_id": ""}, sync=True)
        self.assertFalse(exito)
        self.assertIn("Faltan credenciales", detalle)

    def test_enviar_mensaje_telegram_con_mock_exitoso(self):
        """Valida el envío sincronizado exitoso mockeando la respuesta de urllib."""
        mock_resp = MagicMock()
        mock_resp.status = 200
        mock_resp.__enter__.return_value = mock_resp

        with patch("urllib.request.urlopen", return_value=mock_resp):
            cfg = {"bot_token": "TOK", "chat_id": "123", "activo": True}
            exito, detalle = enviar_mensaje_telegram("Mensaje test", config=cfg, sync=True)
            self.assertTrue(exito)
            self.assertIn("exitosamente", detalle)

    def test_notificadores_telegram_respetan_estado_activo(self):
        """Verifica que los disparadores de alerta respeten el flag activo."""
        with tempfile.TemporaryDirectory() as tmp:
            tmp_root = Path(tmp)
            # Inactivo
            guardar_config_telegram({"bot_token": "T", "chat_id": "C", "activo": False}, tmp_root)
            self.assertFalse(notificar_quiebre_stock_telegram("100", "Prod", 0, 5, tmp_root))
            self.assertFalse(notificar_seguridad_bloqueo_telegram("operario1", root_path=tmp_root))
            self.assertFalse(notificar_cierre_jornada_telegram("Admin", 10, 20, 50, tmp_root))

            # Activo
            guardar_config_telegram({
                "bot_token": "T",
                "chat_id": "C",
                "activo": True,
                "notificar_quiebres": True,
                "notificar_seguridad": True,
                "notificar_cierre": True,
            }, tmp_root)
            with patch("src.telegram_notifier.enviar_mensaje_telegram") as mock_send:
                self.assertTrue(notificar_quiebre_stock_telegram("100", "Prod", 0, 5, tmp_root))
                self.assertTrue(notificar_seguridad_bloqueo_telegram("operario1", root_path=tmp_root))
                self.assertTrue(notificar_cierre_jornada_telegram("Admin", 10, 20, 50, tmp_root))
                self.assertEqual(mock_send.call_count, 3)

    def test_user_store_bloqueo_fuerza_bruta_dispara_alerta_telegram(self):
        """Verifica que 5 intentos fallidos activen el bloqueo y llamen a la alerta de seguridad."""
        with tempfile.TemporaryDirectory() as tmp:
            tmp_root = Path(tmp)
            db_path = tmp_root / "usuarios.sqlite3"
            store = UserStore(db_path)
            store.create_user("usuario_test", "ClaveValida2026!", "operador")

            with patch("src.telegram_notifier.notificar_seguridad_bloqueo_telegram") as mock_notif:
                # 4 intentos fallidos
                for _ in range(4):
                    res = store.authenticate_detailed("usuario_test", "ClaveErronea123!")
                    self.assertEqual(res.status, "invalid")
                mock_notif.assert_not_called()

                # 5to intento fallido -> bloqueo
                res_bloqueo = store.authenticate_detailed("usuario_test", "ClaveErronea123!")
                self.assertEqual(res_bloqueo.status, "locked")
                mock_notif.assert_called_once_with("usuario_test", retry_after_minutos=15)

    def test_sincronizar_catalogo_completo_excel(self):
        """Verifica que sincronizar_catalogo_completo_excel actualice Productos.xlsx con los saldos de SQLite."""
        with tempfile.TemporaryDirectory() as tmp:
            import openpyxl

            tmp_root = Path(tmp)
            ruta_excel = tmp_root / "Productos.xlsx"
            ruta_db = tmp_root / "test_products.sqlite3"

            # Crear libro Excel base (fila 1: encabezados, fila 2: vacía, filas 3+: datos)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Productos"
            ws.append(["Codigo", "Nombre", "Stock", "Precio", "Stock Min", "Stock Max", "Cod Barras"])
            ws.append([])
            ws.append(["101", "Taladro Percutor", 20, 15000, 5, 50, "7791010101"])
            ws.append(["102", "Amoladora Angular", 15, 12000, 3, 30, "7791020202"])
            wb.save(ruta_excel)
            wb.close()

            # Iniciar StockManager aislado con su propia BD
            sm = StockManager(ruta_excel, ruta_base_datos=ruta_db)
            # Modificar en SQLite directamente (simulando descuentos concurrentes masivos)
            sm.product_store.update_stocks({"101": 8, "102": 2})

            # Ejecutar sincronización completa
            filas_sincronizadas = sm.sincronizar_catalogo_completo_excel()
            self.assertEqual(filas_sincronizadas, 2)

            # Verificar que el archivo Excel contenga los nuevos valores
            wb_verif = openpyxl.load_workbook(ruta_excel, data_only=True)
            ws_verif = wb_verif["Productos"]
            stock_101 = ws_verif.cell(3, 3).value
            stock_102 = ws_verif.cell(4, 3).value
            wb_verif.close()

            self.assertEqual(stock_101, 8)
            self.assertEqual(stock_102, 2)


if __name__ == "__main__":
    unittest.main()
