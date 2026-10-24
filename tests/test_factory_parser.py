import tempfile
import unittest
from pathlib import Path
import openpyxl

from src.factory_parser import (
    parsear_archivo_fabrica,
    parsear_pedido_fabrica_excel,
)
from src.stock_manager import StockManager


class FactoryIngestionTests(unittest.TestCase):
    def test_parsear_pedido_fabrica_excel_xlsx(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "pedido_fabrica.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            assert ws is not None
            ws.append(["Código", "Descripción", "Cantidad"])
            ws.append(["101", "Proteína Whey 1kg", 20])
            ws.append(["102", "Creatina 300g", 15])
            ws.append(["101", "Proteína Whey 1kg (Lote B)", 10])
            wb.save(ruta)

            totales, detalle = parsear_pedido_fabrica_excel(ruta)
            self.assertEqual(totales["101"], 30)
            self.assertEqual(totales["102"], 15)
            self.assertEqual(len(detalle), 3)

    def test_stock_manager_registrar_ingreso_fabrica_sumar_y_sobrescribir(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = "Hoja1"
            ws["A1"] = "Codigo"
            ws["B1"] = "Productos"
            ws["C1"] = "Stock"
            ws["A4"] = 101
            ws["B4"] = "Producto A"
            ws["C4"] = 10
            ws["A5"] = 102
            ws["B5"] = "Producto B"
            ws["C5"] = 5
            wb.save(ruta)

            manager = StockManager(ruta)
            self.assertEqual(manager.catalogo["101"][2], 10)
            self.assertEqual(manager.catalogo["102"][2], 5)

            # 1. Sumar
            resumen_suma = manager.registrar_ingreso_fabrica(
                {"101": 15, "102": 5}, sobrescribir=False
            )
            self.assertEqual(resumen_suma["101"], (10, 25))
            self.assertEqual(resumen_suma["102"], (5, 10))
            self.assertEqual(manager.catalogo["101"][2], 25)
            self.assertEqual(manager.catalogo["102"][2], 10)

            # Verificar persistencia en Excel
            wb_verif = openpyxl.load_workbook(ruta, data_only=True)
            self.assertEqual(wb_verif.active["C4"].value, 25)
            self.assertEqual(wb_verif.active["C5"].value, 10)

            # 2. Sobrescribir
            resumen_sobr = manager.registrar_ingreso_fabrica(
                {"101": 50}, sobrescribir=True
            )
            self.assertEqual(resumen_sobr["101"], (25, 50))
            self.assertEqual(manager.catalogo["101"][2], 50)
            wb_verif2 = openpyxl.load_workbook(ruta, data_only=True)
            self.assertEqual(wb_verif2.active["C4"].value, 50)

    def test_registrar_ingreso_producto_inexistente_falla(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = "Hoja1"
            ws["A1"] = "Codigo"
            ws["B1"] = "Productos"
            ws["C1"] = "Stock"
            ws["A4"] = 101
            ws["B4"] = "Producto A"
            ws["C4"] = 10
            wb.save(ruta)

            manager = StockManager(ruta)
            with self.assertRaises(ValueError) as ctx:
                manager.registrar_ingreso_fabrica({"999": 10})
            self.assertIn("999", str(ctx.exception))
