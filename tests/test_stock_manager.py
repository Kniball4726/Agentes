import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import xlrd
from pypdf import PdfReader

from src.consolidator import ItemConsolidado
from src.remito_parser import ItemRemito, ResultadoRemito
from src.stock_manager import RegistroDescuento, StockManager


class StockManagerTests(unittest.TestCase):
    def test_pdf_export_splits_long_catalog_across_a4_pages(self):
        with tempfile.TemporaryDirectory() as directorio:
            libro = openpyxl.Workbook()
            hoja = libro.active
            assert hoja is not None
            for rango in ("A1:A2", "B1:B2", "C1:D2"):
                hoja.merge_cells(rango)
            hoja["A1"] = "CODIGO"
            hoja["B1"] = "PRODUCTOS"
            hoja["C1"] = "FECHA: 22/10/26\nPRODUCTOS: 251\nUNIDADES: 251"
            hoja.append(["CODIGO", "PRODUCTOS", "DESCUENTO"])
            for fila in range(4, 255):
                hoja.append([fila, f"Producto {fila}", None])
            ruta_pdf = Path(directorio) / "informe.pdf"

            StockManager._guardar_informe_pdf(libro, ruta_pdf)

            reader = PdfReader(ruta_pdf)
            self.assertGreater(len(reader.pages), 1)
            texto = "\n".join(
                pagina.extract_text() or "" for pagina in reader.pages
            )
            self.assertIn("Producto 254", texto)
            self.assertIn("CODIGO", texto)
            self.assertNotIn("\nUNIDADES\n", texto)

    def test_imports_duplicate_catalog_rows_and_carries_stock_forward_in_sqlite(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja.title = "Hoja1"
            hoja["A1"] = "Codigo"
            hoja["C1"] = "Productos"
            hoja["C3"] = "Línea"
            hoja["D3"] = "Stock Inicial"
            hoja["E3"] = "Descuento"
            hoja["F3"] = "Stock Final"
            hoja["A4"] = 100
            hoja["C4"] = "Producto duplicado"
            hoja["A5"] = 100
            hoja["C5"] = "Producto duplicado"
            hoja["A6"] = 200
            hoja["C6"] = "Producto con stock"
            hoja["D6"] = 7
            libro.save(ruta)

            manager = StockManager(ruta)

            self.assertEqual(len(manager.productos_base), 2)
            self.assertEqual(manager.catalogo["100"][2], 0)
            self.assertEqual(manager.catalogo["200"][2], 7)

            manager.actualizar_stock_base([
                RegistroDescuento(
                    codigo="100",
                    nombre_producto="Producto duplicado",
                    fila_excel=3,
                    stock_anterior=0,
                    cantidad_descontada=2,
                    stock_nuevo=-2,
                )
            ])

            libro_plantilla = openpyxl.load_workbook(ruta, data_only=True)
            hoja_plantilla = libro_plantilla["Hoja1"]
            self.assertIsNone(hoja_plantilla.cell(4, 4).value)
            self.assertIsNone(hoja_plantilla.cell(4, 6).value)

            manager_siguiente_corrida = StockManager(ruta)
            self.assertEqual(manager_siguiente_corrida.catalogo["100"][2], -2)
            self.assertEqual(manager_siguiente_corrida.catalogo["200"][2], 7)
            self.assertEqual(len(manager_siguiente_corrida.filas_catalogo), 3)
            self.assertEqual(
                [fila[1] for fila in manager_siguiente_corrida.filas_catalogo],
                ["100", "100", "200"],
            )

    def test_keeps_legacy_a_to_c_catalog_layout(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            hoja.append([])
            hoja.append([345, "Producto clásico", 12])
            libro.save(ruta)

            manager = StockManager(ruta)

            self.assertEqual(manager.productos_base, [(3, "345", "Producto clásico", 12)])
            self.assertEqual(manager._col_stock, 3)
            self.assertIsNone(manager._col_stock_final)

    def test_adds_product_to_sqlite_and_template_in_reference_order(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja["A1"] = "Codigo"
            hoja["C1"] = "Productos"
            hoja["D3"] = "Stock Inicial"
            hoja["E3"] = "Descuento"
            hoja["F3"] = "Stock Final"
            hoja["A4"] = 100
            hoja["C4"] = "Producto existente"
            hoja["A4"].fill = openpyxl.styles.PatternFill(
                "solid", fgColor="CCCCFF"
            )
            libro.save(ruta)

            manager = StockManager(ruta)
            contenido_plantilla = ruta.read_bytes()
            exito, _, fila = manager.agregar_producto_en_posicion(
                codigo_nuevo="101",
                nombre_nuevo="Producto nuevo",
                codigo_arriba="100",
                stock_inicial=5,
            )

            self.assertTrue(exito)
            self.assertEqual(fila, 5)
            self.assertNotEqual(ruta.read_bytes(), contenido_plantilla)
            respaldo = ruta.with_name("Productos_backup.xlsx")
            self.assertEqual(respaldo.read_bytes(), contenido_plantilla)
            libro_actualizado = openpyxl.load_workbook(ruta, data_only=True)
            hoja_actualizada = libro_actualizado.active
            self.assertEqual(hoja_actualizada["A5"].value, 101)
            self.assertEqual(hoja_actualizada["C5"].value, "Producto nuevo")
            self.assertEqual(hoja_actualizada["D5"].value, 5)
            self.assertEqual(
                hoja_actualizada["A5"].fill.fgColor.rgb,
                hoja_actualizada["A4"].fill.fgColor.rgb,
            )
            nuevo_manager = StockManager(ruta)
            self.assertEqual(nuevo_manager.catalogo["101"][2], 5)
            self.assertEqual(
                [producto[1] for producto in nuevo_manager.productos_base],
                ["100", "101"],
            )
            self.assertEqual(nuevo_manager.filas_catalogo[1][0], 5)
            with patch.object(
                StockManager,
                "_guardar_informe_xls",
                wraps=StockManager._guardar_informe_xls,
            ) as guardar:
                nuevo_manager.aplicar_descuentos(
                    {},
                    Path(directorio) / "informes",
                )
            hoja_reporte = guardar.call_args.args[0].active
            self.assertEqual(hoja_reporte["B4"].value, "Producto existente")
            self.assertEqual(hoja_reporte["B5"].value, "Producto nuevo")

    def test_edits_and_deactivates_product_in_catalog_and_template(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja["A1"] = "Codigo"
            hoja["C1"] = "Productos"
            hoja["D3"] = "Stock Inicial"
            hoja["E3"] = "Descuento"
            hoja["F3"] = "Stock Final"
            for fila in (4, 5):
                hoja.cell(fila, 1, 100)
                hoja.cell(fila, 3, "Producto repetido")
                hoja.cell(fila, 4, 8)
            hoja["A6"] = 200
            hoja["C6"] = "Otro producto"
            hoja["D6"] = 3
            libro.save(ruta)
            contenido_original = ruta.read_bytes()

            manager = StockManager(ruta)
            manager.actualizar_producto("100", "101", "Producto editado", 12)

            respaldo = ruta.with_name("Productos_backup.xlsx")
            self.assertEqual(respaldo.read_bytes(), contenido_original)
            actualizado = openpyxl.load_workbook(ruta, data_only=True)
            hoja_actualizada = actualizado.active
            for fila in (4, 5):
                self.assertEqual(hoja_actualizada.cell(fila, 1).value, 101)
                self.assertEqual(
                    hoja_actualizada.cell(fila, 3).value, "Producto editado"
                )
                self.assertEqual(hoja_actualizada.cell(fila, 4).value, 12)
                self.assertIsNone(hoja_actualizada.cell(fila, 5).value)
                self.assertEqual(hoja_actualizada.cell(fila, 6).value, 12)

            self.assertNotIn("100", manager.catalogo)
            self.assertEqual(manager.catalogo["101"][2], 12)
            manager.cambiar_estado_producto("101", False)
            self.assertNotIn("101", manager.catalogo)
            self.assertEqual(
                [producto["code"] for producto in manager.listar_productos()],
                ["101", "200"],
            )
            with patch.object(
                StockManager,
                "_guardar_informe_xls",
                wraps=StockManager._guardar_informe_xls,
            ) as guardar:
                manager.aplicar_descuentos(
                    {},
                    Path(directorio) / "informes",
                )
            hoja_reporte = guardar.call_args.args[0].active
            self.assertEqual(hoja_reporte["A4"].value, 200)
            self.assertEqual(hoja_reporte["B4"].value, "Otro producto")
            manager.cambiar_estado_producto("101", True)
            self.assertIn("101", manager.catalogo)

    def test_builds_discount_report_from_read_only_template(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "Productos.xlsx"
            libro = openpyxl.Workbook()
            hoja = libro.active
            hoja.merge_cells("A1:A2")
            hoja.merge_cells("B1:B2")
            hoja.merge_cells("C1:C2")
            hoja.merge_cells("D1:F2")
            hoja["A1"] = "Codigo"
            hoja["B1"] = "Productos"
            hoja["C1"] = "Descuento"
            hoja["B3"] = "Línea"
            hoja["C3"] = "Stock Inicial"
            hoja["D3"] = "Descuento"
            hoja["E3"] = "Stock Final"
            hoja["A4"] = 101
            hoja["B4"] = "Producto con descuento"
            hoja["C4"] = 10
            hoja["A5"] = 202
            hoja["B5"] = "Producto sin descuento"
            hoja["C5"] = 3
            hoja["A4"].fill = openpyxl.styles.PatternFill("solid", fgColor="CCCCFF")
            hoja["D4"].fill = openpyxl.styles.PatternFill("solid", fgColor="FF8080")
            hoja.column_dimensions["B"].width = 42
            libro.save(ruta)
            contenido_plantilla = ruta.read_bytes()

            manager = StockManager(ruta)
            pedido = ResultadoRemito(
                archivo="pedido.pdf",
                numero_remito="R-123",
                fecha="Hoy",
                cliente="Cliente Ejemplo",
                articulos_declarados=3,
                articulos_extraidos=1,
                verificacion_ok=False,
                items=[ItemRemito("101", "Producto con descuento", 2, "R-123")],
                descripciones_sin_coincidencia=[
                    ("Producto sin coincidencia", "No se encontró en el catálogo")
                ],
            )
            with patch.object(StockManager, "_guardar_informe_xls", wraps=StockManager._guardar_informe_xls) as guardar:
                reporte = manager.aplicar_descuentos(
                    {
                        "101": ItemConsolidado("101", "Producto con descuento", 2),
                        "999": ItemConsolidado("999", "Producto no catalogado", 1),
                    },
                    Path(directorio) / "informes",
                    pedidos_procesados=[pedido],
                )

            libro_salida = guardar.call_args.args[0]
            hoja_reporte = libro_salida.active
            self.assertEqual(
                {hoja_reporte.cell(1, columna).font.sz for columna in range(1, 4)},
                {9},
            )
            self.assertEqual(
                {hoja_reporte.cell(3, columna).font.sz for columna in range(1, 4)},
                {9},
            )
            ruta_xls = Path(directorio) / "informe.xls"
            StockManager._guardar_informe_xls(libro_salida, ruta_xls)
            libro_xls = xlrd.open_workbook(str(ruta_xls), formatting_info=True)
            hoja_xls = libro_xls.sheet_by_index(0)
            for columna in range(4):
                indice_xf = hoja_xls.cell_xf_index(0, columna)
                indice_fuente = libro_xls.xf_list[indice_xf].font_index
                self.assertEqual(libro_xls.font_list[indice_fuente].height, 180)
            archivos_pdf = list((Path(directorio) / "informes").glob("*.pdf"))
            self.assertEqual(len(archivos_pdf), 1)
            self.assertEqual(reporte.archivo_pdf, str(archivos_pdf[0]))
            reader_pdf = PdfReader(archivos_pdf[0])
            texto_pdf = "\n".join(
                pagina.extract_text() or "" for pagina in reader_pdf.pages
            )
            self.assertIn("PRODUCTO CON DESCUENTO", texto_pdf.upper())
            self.assertIn("TOTAL UNIDADES DESCONTADAS", texto_pdf.upper())
            self.assertEqual(
                round(float(reader_pdf.pages[0].mediabox.width)),
                595,
            )
            self.assertRegex(
                hoja_reporte["C1"].value,
                r"^FECHA: \d{2}/\d{2}/\d{2}$",
            )
            self.assertNotIn("PRODUCTOS:", hoja_reporte["C1"].value)
            self.assertNotIn("UNIDADES:", hoja_reporte["C1"].value)
            self.assertEqual(hoja_reporte["A1"].value, "Codigo")
            self.assertNotEqual(hoja_reporte["B1"].value, "UNIDADES")
            self.assertEqual(hoja_reporte["B3"].value, "Línea")
            self.assertEqual(hoja_reporte["C3"].value, "DESCUENTO")
            self.assertEqual(
                [hoja_reporte.cell(4, columna).value for columna in range(1, 4)],
                [101, "Producto con descuento", -2],
            )
            self.assertEqual(
                [hoja_reporte.cell(5, columna).value for columna in range(1, 4)],
                [202, "Producto sin descuento", None],
            )
            self.assertEqual(
                hoja_reporte["B6"].value,
                "TOTAL CÓDIGOS DE PRODUCTO DESCONTADOS:",
            )
            self.assertEqual(hoja_reporte["C6"].value, 1)
            self.assertEqual(hoja_reporte["B7"].value, "TOTAL UNIDADES DESCONTADAS:")
            self.assertEqual(hoja_reporte["C7"].value, 2)
            self.assertEqual(hoja_reporte["A1"].fill.fgColor.rgb, "00D9EAF7")
            self.assertEqual(hoja_reporte["A1"].font.color.rgb, "0017365D")
            self.assertTrue(hoja_reporte["A1"].font.bold)
            self.assertGreaterEqual(hoja_reporte.column_dimensions["C"].width, 22)
            self.assertEqual(hoja_reporte["A3"].fill.fgColor.rgb, "00D9EAF7")
            self.assertEqual(hoja_reporte["A8"].fill.fgColor.rgb, "00EAF0F6")
            self.assertEqual(hoja_reporte["A9"].fill.fgColor.rgb, "00D9EAF7")
            self.assertEqual(hoja_reporte.max_column, 4)
            self.assertIn(
                "C1:D2",
                {str(rango) for rango in hoja_reporte.merged_cells.ranges},
            )
            valores_informe = [
                celda.value
                for fila in hoja_reporte.iter_rows()
                for celda in fila
                if celda.value is not None
            ]
            self.assertIn("PEDIDOS CON PRODUCTOS DESCONTADOS", valores_informe)
            self.assertIn("Cliente Ejemplo", valores_informe)
            self.assertIn("R-123", valores_informe)
            self.assertIn("PRODUCTOS NO DESCONTADOS - SIN COINCIDENCIA", valores_informe)
            self.assertIn("Producto sin coincidencia", valores_informe)
            self.assertIn("Producto no catalogado", valores_informe)
            self.assertEqual(hoja_reporte.freeze_panes, "A4")
            self.assertFalse(hoja_reporte.column_dimensions["B"].hidden)
            self.assertEqual(hoja_reporte["C4"].fill.fgColor.rgb, "00FF8080")
            self.assertFalse(any("STOCK INICIAL" in str(valor).upper() for valor in valores_informe))
            self.assertFalse(any("STOCK FINAL" in str(valor).upper() for valor in valores_informe))
            self.assertEqual(
                len([fila for fila in range(4, hoja_reporte.max_row + 1)
                     if hoja_reporte.cell(fila, 1).value in (101, 202)]),
                2,
            )
            self.assertEqual(str(hoja_reporte.page_setup.paperSize), "9")
            self.assertEqual(hoja_reporte.page_setup.orientation, "portrait")
            self.assertEqual(hoja_reporte.page_setup.fitToWidth, 1)
            for lado in ("left", "right", "top", "bottom"):
                self.assertAlmostEqual(
                    getattr(hoja_reporte.page_margins, lado) * 2.54,
                    2,
                )
            self.assertTrue(Path(guardar.call_args.args[1]).exists())
            self.assertEqual(ruta.read_bytes(), contenido_plantilla)

    def test_actualizar_stock_base_records_dispatch_and_audit(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            db_path = Path(temp_dir) / "productos.sqlite3"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["1001", "Item A", 20])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=db_path)
            descuentos = [
                RegistroDescuento(
                    codigo="1001",
                    nombre_producto="Item A",
                    fila_excel=3,
                    stock_anterior=20,
                    cantidad_descontada=5,
                    stock_nuevo=15,
                )
            ]
            actualizados = manager.actualizar_stock_base(
                descuentos=descuentos,
                sincronizar_excel=False,
                nombre_archivo="informe_test.xlsx",
            )
            self.assertEqual(actualizados, 1)

            prod_map = {p["code"]: p["stock"] for p in manager.product_store.list_products()}
            self.assertIn("1001", prod_map)
            self.assertEqual(prod_map["1001"], 15)

            conn = manager.product_store._connect()
            try:
                despacho = conn.execute(
                    "SELECT * FROM despachos WHERE nombre_archivo = 'informe_test.xlsx'"
                ).fetchone()
                self.assertIsNotNone(despacho)
                self.assertEqual(despacho["total_articulos"], 5)

                detalle = conn.execute(
                    "SELECT * FROM despacho_detalles WHERE product_code = '1001'"
                ).fetchone()
                self.assertIsNotNone(detalle)
                self.assertEqual(detalle["cantidad_descontada"], 5)
                self.assertEqual(detalle["stock_nuevo"], 15)
            finally:
                conn.close()

    def test_generar_pedido_sugerido_fabrica_excel(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            db_path = Path(temp_dir) / "productos.sqlite3"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["P1", "Item Agotado", -2])
            ws.append(["P2", "Item Quiebre", 3])
            ws.append(["P3", "Item Optimo", 50])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=db_path)
            ruta_generada = manager.generar_pedido_sugerido_fabrica(
                ruta_salida=Path(temp_dir) / "Pedido_Test.xlsx",
                solo_criticos=True,
            )
            self.assertTrue(ruta_generada.exists())

            wb_out = openpyxl.load_workbook(ruta_generada)
            ws_out = wb_out["Pedido a Fabrica"]
            self.assertEqual(ws_out["A1"].value, "SUGERENCIA DE PEDIDO A FÁBRICA")

            headers = [ws_out.cell(4, c).value for c in range(1, 8)]
            self.assertEqual(
                headers,
                [
                    "CÓDIGO",
                    "PRODUCTO",
                    "STOCK ACTUAL",
                    "STOCK MÍNIMO",
                    "STOCK MÁXIMO",
                    "CANTIDAD SUGERIDA A PEDIR",
                    "ESTADO",
                ],
            )

            codigos_pedidos = [ws_out.cell(row, 1).value for row in (5, 6)]
            self.assertIn("P1", codigos_pedidos)
            self.assertIn("P2", codigos_pedidos)
            self.assertNotIn("P3", codigos_pedidos)

            row_p1 = 5 if ws_out.cell(5, 1).value == "P1" else 6
            self.assertEqual(ws_out.cell(row_p1, 3).value, -2)
            self.assertEqual(ws_out.cell(row_p1, 4).value, 5)
            self.assertEqual(ws_out.cell(row_p1, 5).value, 100)
            self.assertEqual(ws_out.cell(row_p1, 6).value, 102)
            self.assertEqual(ws_out.cell(row_p1, 7).value, "AGOTADO")

            row_p2 = 6 if row_p1 == 5 else 5
            self.assertEqual(ws_out.cell(row_p2, 3).value, 3)
            self.assertEqual(ws_out.cell(row_p2, 4).value, 5)
            self.assertEqual(ws_out.cell(row_p2, 5).value, 100)
            self.assertEqual(ws_out.cell(row_p2, 6).value, 97)
            self.assertEqual(ws_out.cell(row_p2, 7).value, "QUIEBRE_INMINENTE")

            wb_out.close()

    def test_registrar_ingreso_fabrica_permite_negativos(self):
        """Descuento manual con valor negativo debe reducir el stock, incluyendo quedar negativo."""
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["P1", "Producto A", 10])
            ws.append(["P2", "Producto B", 3])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=Path(temp_dir) / "db.sqlite3")

            # Descuento parcial: 10 + (-4) = 6
            res = manager.registrar_ingreso_fabrica({"P1": -4}, sobrescribir=False)
            self.assertEqual(res["P1"], (10, 6))

            # Descuento que deja en negativo: 3 + (-5) = -2
            manager2 = StockManager(excel_path, ruta_base_datos=Path(temp_dir) / "db.sqlite3")
            res2 = manager2.registrar_ingreso_fabrica({"P2": -5}, sobrescribir=False)
            self.assertEqual(res2["P2"], (3, -2))

            # Sobrescribir nunca queda negativo (clamp a 0)
            manager3 = StockManager(excel_path, ruta_base_datos=Path(temp_dir) / "db.sqlite3")
            res3 = manager3.registrar_ingreso_fabrica({"P1": -99}, sobrescribir=True)
            self.assertEqual(res3["P1"][1], 0)

    def test_descontar_mercado_full(self):
        """Descontar mercado full reduce el stock y registra el despacho."""
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            db_path = Path(temp_dir) / "db.sqlite3"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["F1", "Filtro Aceite", 20])
            ws.append(["F2", "Bujía NGK", 15])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=db_path)
            resumen = manager.descontar_mercado_full({"F1": 5, "F2": 3}, nombre_archivo="Envio_Full_01.xlsx")

            self.assertEqual(resumen["F1"], (20, 15))
            self.assertEqual(resumen["F2"], (15, 12))

            # Verificar que quedó registrado en SQLite
            manager2 = StockManager(excel_path, ruta_base_datos=db_path)
            self.assertEqual(manager2.catalogo["F1"][2], 15)
            self.assertEqual(manager2.catalogo["F2"][2], 12)

    def test_actualizar_stock_base_crea_notificacion_para_operario(self):
        """Cuando un operario descuenta stock, se genera una notificación para admin."""
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            db_path = Path(temp_dir) / "db.sqlite3"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["OP1", "Producto Operario", 50])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=db_path)
            descuentos = [
                RegistroDescuento("OP1", "Producto Operario", 50, 10, 40, 4, False)
            ]
            manager.actualizar_stock_base(
                descuentos=descuentos,
                nombre_archivo="remito_operador.pdf",
                usuario="carlos_operario",
            )

            self.assertEqual(manager.contar_notificaciones_no_leidas(), 1)
            notifs = manager.obtener_notificaciones_admin()
            self.assertEqual(len(notifs), 1)
            self.assertEqual(notifs[0]["operator_username"], "carlos_operario")
            self.assertIn("carlos_operario", notifs[0]["message"])

    def test_procesar_stock_fisico_genera_comparativa_excel_y_notificacion(self):
        """Relevar stock físico actualiza stock_fisico, compara con digital, genera Excel unificado y notifica al admin."""
        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "Productos.xlsx"
            db_path = Path(temp_dir) / "db.sqlite3"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
            ws.append([])
            ws.append(["A1", "Producto Exacto", 10])
            ws.append(["A2", "Producto Faltante", 20])
            ws.append(["A3", "Producto Sobrante", 5])
            wb.save(excel_path)

            manager = StockManager(excel_path, ruta_base_datos=db_path)
            ruta_salida = Path(temp_dir) / "Comparativa_Unificada.xlsx"
            ruta_doc, comparativa = manager.procesar_stock_fisico(
                {"A1": 10, "A2": 15, "A3": 8},
                usuario="lucia_operaria",
                nombre_archivo="Conteo_Lunes.xlsx",
                ruta_salida=ruta_salida,
            )

            self.assertTrue(ruta_doc.exists())
            mapa = {item["code"]: item for item in comparativa}
            self.assertEqual(mapa["A1"]["diferencia"], 0)
            self.assertEqual(mapa["A1"]["estado"], "COINCIDE")
            self.assertEqual(mapa["A2"]["diferencia"], -5)
            self.assertEqual(mapa["A2"]["estado"], "FALTANTE")
            self.assertEqual(mapa["A3"]["diferencia"], 3)
            self.assertEqual(mapa["A3"]["estado"], "SOBRANTE")

            # Verificar que se creó la notificación para el administrador
            notifs = manager.obtener_notificaciones_admin()
            self.assertEqual(len(notifs), 1)
            self.assertEqual(notifs[0]["operator_username"], "lucia_operaria")
            self.assertEqual(notifs[0]["notification_type"], "stock_fisico")
            self.assertIn("Comparativa_Unificada.xlsx", notifs[0]["message"])

            # Verificar que el Excel unificado tiene las 2 hojas
            wb_comp = openpyxl.load_workbook(ruta_doc)
            self.assertIn("Comparativa y Diferencias", wb_comp.sheetnames)
            self.assertIn("Catálogo Unificado", wb_comp.sheetnames)
            wb_comp.close()


if __name__ == "__main__":
    unittest.main()