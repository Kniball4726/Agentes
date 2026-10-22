import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.mercado_libre_parser import _buscar_codigo, parsear_pedido_mercado_libre


class MercadoLibreParserTests(unittest.TestCase):
    def test_only_numbered_product_lines_are_matched_and_quantity_is_applied(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "pedido.docx"
            documento = Document()
            documento.add_paragraph(
                "2 Creatina 300 g UT\n"
                "Cliente: Creatina 300 g UT\n"
                "1. BCAA 120 cápsulas UT\n"
                "Entrega por la tarde"
            )
            documento.add_paragraph("Creatina 300 g UT")
            tabla = documento.add_table(rows=1, cols=2)
            tabla.cell(0, 0).text = "3"
            tabla.cell(0, 1).text = "Creatina 300 g UT"
            documento.add_paragraph("1 Producto inexistente")
            documento.save(ruta)

            resultado = parsear_pedido_mercado_libre(
                ruta,
                [
                    ("100", "Creatina 300 g UT"),
                    ("200", "BCAA 120 cápsulas UT"),
                ],
            )

            self.assertEqual(resultado.articulos_declarados, 7)
            self.assertEqual(resultado.articulos_extraidos, 6)
            self.assertEqual(
                [(item.codigo, item.cantidad) for item in resultado.items],
                [("100", 2), ("200", 1), ("100", 3)],
            )
            self.assertEqual(
                resultado.descripciones_sin_coincidencia,
                [
                    (
                        "1 Producto inexistente",
                        "No se encontró una coincidencia suficientemente clara en el catálogo",
                    )
                ],
            )
            self.assertFalse(resultado.verificacion_ok)

    def test_groups_only_unmatched_products_with_the_same_description(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "pedido.docx"
            documento = Document()
            documento.add_paragraph("2 Whey Choco")
            documento.add_paragraph("2 whey choco")
            documento.add_paragraph("1 Creatina 300 g UT")
            documento.add_paragraph("3 Creatina 300 g UT")
            documento.save(ruta)

            resultado = parsear_pedido_mercado_libre(
                ruta,
                [("100", "Creatina 300 g UT")],
            )

            self.assertEqual(resultado.articulos_declarados, 8)
            self.assertEqual(resultado.articulos_extraidos, 4)
            self.assertEqual(
                [(item.codigo, item.cantidad) for item in resultado.items],
                [("100", 1), ("100", 3)],
            )
            self.assertEqual(
                len(resultado.descripciones_sin_coincidencia),
                1,
            )
            self.assertEqual(
                resultado.descripciones_sin_coincidencia[0][0],
                "4 Whey Choco",
            )
            self.assertFalse(resultado.verificacion_ok)

    def test_matches_unambiguous_approximate_description(self):
        codigo, motivo = _buscar_codigo(
            "Garcinia Cambogia x 60 comp Vitatech",
            {"209": ["Garcinia Gambogia Plus x 60 comp. VT"]},
        )

        self.assertEqual(codigo, "209")
        self.assertEqual(motivo, "")

    def test_matches_catalog_products_with_additional_catalog_details(self):
        casos = [
            (
                "CREATINA + ELECTROLYTES ULTRA TECH",
                {"2763": ["Creatine + ELECTROLYTES x 312 grs"]},
                "2763",
            ),
            (
                "SABORIZADA CREATINA 300 G ULTRA TECH",
                {"2753": ["SABORIZADA CREATINA 300 G ULTRA TECH FRUTOS ROJOS C/ LIMON"]},
                "2753",
            ),
            (
                "ULTRA TECH MULTI COLLAGEN",
                {"2555": ["MULTI COLLAGEN X 405 GR ULTRA TECH"]},
                "2555",
            ),
            (
                "ULTRA TECH BCAA 2000",
                {"28": ["ULTRA TECH AMINOACIDOS BCAA 2000"]},
                "28",
            ),
        ]

        for descripcion, catalogo, codigo_esperado in casos:
            with self.subTest(descripcion=descripcion):
                codigo, motivo = _buscar_codigo(descripcion, catalogo)

                self.assertEqual(codigo, codigo_esperado)
                self.assertEqual(motivo, "")

    def test_does_not_match_when_multiple_catalog_products_contain_the_terms(self):
        codigo, motivo = _buscar_codigo(
            "ULTRA TECH BCAA 2000",
            {
                "28": ["ULTRA TECH AMINOACIDOS BCAA 2000 VAINILLA"],
                "29": ["ULTRA TECH AMINOACIDOS BCAA 2000 CHOCOLATE"],
            },
        )

        self.assertIsNone(codigo)
        self.assertIn("ambigua", motivo)

    def test_does_not_match_a_missing_whey_flavor(self):
        codigo, _ = _buscar_codigo(
            "WHEY PROTEIN X 907 GR ULTRA TECH COCO",
            {
                "2244": ["WHEY PROTEIN X 907 GR ULTRA TECH CHOCOLATE"],
                "2245": ["WHEY PROTEIN X 907 GR ULTRA TECH VAINILLA"],
            },
        )

        self.assertIsNone(codigo)

    def test_does_not_guess_between_similar_catalog_products(self):
        codigo, motivo = _buscar_codigo(
            "BISGLICINATO DE MAGNESIO X 60 CAPS VITATECH",
            {
                "2724": ["BISGLICINATO DE MAGNESIO X 60 CAPS"],
                "2798": ["BISGLICINATO MAGNESIO * 60 CAPS UT"],
            },
        )

        self.assertIsNone(codigo)
        self.assertIn("ambigua", motivo)

    def test_does_not_match_similar_but_different_product(self):
        codigo, _ = _buscar_codigo(
            "ZINC X 30 COMP VITATECH",
            {"116": ["Vision x 30 comp VT"]},
        )

        self.assertIsNone(codigo)


if __name__ == "__main__":
    unittest.main()
