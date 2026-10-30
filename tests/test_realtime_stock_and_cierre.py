"""Pruebas unitarias para el descuento en tiempo real en SQLite y Productos.xlsx,
restitución/regresión de stock al anular o desarmar pedidos y reporte exacto en Cierre de Jornada.
"""
from pathlib import Path
import tempfile
import openpyxl

from src.consolidator import ItemConsolidado
from src.orders_store import OrderStore, Pedido
from src.product_store import ProductStore
from src.stock_manager import StockManager, RegistroDescuento
from src.ui.dialogs.dlg_pedidos import _restituir_stock_pedido


def test_descuento_real_time_y_restitucion_stock(tmp_path: Path):
    """Verifica que el stock se actualice en SQLite y Productos.xlsx, y que se restituya al desarmar."""
    excel_path = tmp_path / "Productos.xlsx"
    db_prod_path = tmp_path / "productos.sqlite3"
    db_ped_path = tmp_path / "pedidos.sqlite3"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
    ws.append([])
    ws.append(["PROD-1", "Aceite Motor", 50])
    ws.append(["PROD-2", "Filtro Aire", 20])
    wb.save(excel_path)
    wb.close()

    stock_mgr = StockManager(excel_path, ruta_base_datos=db_prod_path)
    order_store = OrderStore(db_ped_path)

    # 1. Simular descuento en tiempo real al agregar pedido
    descuentos = [
        RegistroDescuento(
            codigo="PROD-1",
            nombre_producto="Aceite Motor",
            fila_excel=3,
            stock_anterior=50,
            cantidad_descontada=5,
            stock_nuevo=45,
        ),
        RegistroDescuento(
            codigo="PROD-2",
            nombre_producto="Filtro Aire",
            fila_excel=4,
            stock_anterior=20,
            cantidad_descontada=3,
            stock_nuevo=17,
        ),
    ]

    # Guardar en SQLite y Productos.xlsx (sincronizar_excel=True)
    stock_mgr.actualizar_stock_base(
        descuentos=descuentos,
        sincronizar_excel=True,
        nombre_archivo="Remito_101.pdf",
        usuario="Operario1",
    )

    # Verificar que SQLite tenga el stock descontado
    assert stock_mgr.catalogo["PROD-1"][2] == 45
    assert stock_mgr.catalogo["PROD-2"][2] == 17

    # Verificar que Productos.xlsx tenga el stock descontado
    wb_read = openpyxl.load_workbook(excel_path, data_only=True)
    ws_read = wb_read.active
    assert ws_read.cell(3, 3).value == 45
    assert ws_read.cell(4, 3).value == 17
    wb_read.close()

    # 2. Registrar el pedido con items en OrderStore
    items_descontados = [
        ("PROD-1", "Aceite Motor", 5),
        ("PROD-2", "Filtro Aire", 3),
    ]
    pedido = order_store.registrar_pedido(
        nota_pedido="REM-101",
        numero_pedido="1",
        cliente="Cliente Test",
        bultos=2,
        logistica="Domicilio",
        armador="Operario1",
        tipo_archivo="pdf",
        nombre_archivo="Remito_101.pdf",
        items=items_descontados,
    )

    # 3. Probar restitución de stock al anular/desarmar pedido
    class DummyInterfaz:
        def __init__(self, ost, ruta_excel, ruta_db):
            self.order_store = ost
            self.ruta_base_datos = ruta_db
            self.menu = type("DummyMenu", (), {"ruta_excel_base": ruta_excel, "ruta_base_datos": ruta_db})()

    # Ejecutar restitución usando el helper
    interfaz_dummy = DummyInterfaz(order_store, excel_path, db_prod_path)
    unidades_devueltas = _restituir_stock_pedido(interfaz_dummy, pedido)
    assert unidades_devueltas == 8

    # Verificar que el stock volvió a los valores originales (50 y 20)
    stock_mgr_after = StockManager(excel_path, ruta_base_datos=db_prod_path)
    assert stock_mgr_after.catalogo["PROD-1"][2] == 50
    assert stock_mgr_after.catalogo["PROD-2"][2] == 20

    wb_after = openpyxl.load_workbook(excel_path, data_only=True)
    ws_after = wb_after.active
    assert ws_after.cell(3, 3).value == 50
    assert ws_after.cell(4, 3).value == 20
    wb_after.close()


def test_cierre_jornada_descuento_ya_aplicado(tmp_path: Path):
    """Verifica que el informe de Cierre con descuento_ya_aplicado=True no descuente dos veces en memoria."""
    excel_path = tmp_path / "Productos.xlsx"
    db_prod_path = tmp_path / "productos.sqlite3"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CÓDIGO", "PRODUCTOS", "STOCK"])
    ws.append([])
    ws.append(["P-100", "Batería 12V", 3])  # Stock actual ya descontado en tiempo real
    wb.save(excel_path)
    wb.close()

    stock_mgr = StockManager(excel_path, ruta_base_datos=db_prod_path)

    items_consolidados = {
        "P-100": ItemConsolidado(
            codigo="P-100",
            descripcion="Batería 12V",
            cantidad_total=2,
            remitos_origen=["REM-01", "REM-02"],
        )
    }

    carpeta_descuentos = tmp_path / "informes"
    carpeta_descuentos.mkdir()

    # Generar informe con descuento_ya_aplicado=True
    reporte = stock_mgr.aplicar_descuentos(
        items_consolidados=items_consolidados,
        carpeta_destino=carpeta_descuentos,
        guardar_archivo=True,
        prefijo_archivo="Cierre_Descuento",
        descuento_ya_aplicado=True,
    )

    reg = reporte.descuentos_aplicados[0]
    # En tiempo real el stock actual es 3. El pedido descontó 2.
    # El stock inicial era 5, y el nuevo es 3 (no 1).
    assert reg.cantidad_descontada == 2
    assert reg.stock_anterior == 5
    assert reg.stock_nuevo == 3

    # No debe considerarse en quiebre (stock_nuevo es 3 > 0)
    assert len(reporte.articulos_quiebre) == 0
    # Es crítico porque está entre 1 y 5
    assert len(reporte.articulos_criticos) == 1
    assert reporte.articulos_criticos[0].stock_nuevo == 3


def test_obtener_items_pendientes_cierre_deduplicacion(tmp_path: Path):
    """Verifica que obtener_items_pendientes_cierre agrupe por código y consolide remitos."""
    db_ped_path = tmp_path / "pedidos.sqlite3"
    order_store = OrderStore(db_ped_path)

    # Pedido 1
    p1 = order_store.registrar_pedido(
        nota_pedido="REM-1",
        numero_pedido="1",
        cliente="Cliente A",
        bultos=1,
        logistica="Domicilio",
        armador="Op1",
        tipo_archivo="pdf",
        nombre_archivo="doc1.pdf",
        items=[("COD1", "Producto Uno (Desc A)", 3), ("COD2", "Producto Dos", 2)],
    )

    # Pedido 2 con el mismo COD1 pero descripción ligeramente distinta
    p2 = order_store.registrar_pedido(
        nota_pedido="REM-2",
        numero_pedido="2",
        cliente="Cliente B",
        bultos=1,
        logistica="Canje",
        armador="Op1",
        tipo_archivo="excel",
        nombre_archivo="doc2.xlsx",
        items=[("COD1", "Producto Uno (Desc B)", 4)],
    )

    pendientes = order_store.obtener_items_pendientes_cierre()
    mapa = {cod: (cant, remitos) for cod, _, cant, remitos in pendientes}

    assert "COD1" in mapa
    assert mapa["COD1"][0] == 7  # 3 + 4
    assert set(mapa["COD1"][1]) == {"REM-1", "REM-2"}

    assert "COD2" in mapa
    assert mapa["COD2"][0] == 2
    assert mapa["COD2"][1] == ["REM-1"]
