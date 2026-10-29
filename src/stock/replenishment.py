"""Análisis y generación de planillas de sugerencia de reposición a fábrica."""

from datetime import datetime
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generar_pedido_sugerido_fabrica_excel(
    analisis: List,
    ruta_excel_base: Path,
    ruta_salida: Path | str | None = None,
) -> Path:
    """Genera un archivo Excel formateado con el pedido sugerido a fábrica."""
    items_a_pedir = [p for p in analisis if int(p["cantidad_sugerida_pedir"]) > 0]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedido a Fabrica"

    fill_titulo = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_titulo = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_datos = Font(name="Segoe UI", size=10)
    font_destacado = Font(name="Segoe UI", size=10, bold=True, color="0F766E")
    font_agotado = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    font_quiebre = Font(name="Segoe UI", size=10, bold=True, color="D97706")

    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    borde_totales = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("A1:G1")
    celda_titulo = ws["A1"]
    celda_titulo.value = "SUGERENCIA DE PEDIDO A FÁBRICA"
    celda_titulo.fill = fill_titulo
    celda_titulo.font = font_titulo
    celda_titulo.alignment = align_center
    ws.row_dimensions[1].height = 36

    fecha_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ws.merge_cells("A2:G2")
    celda_meta = ws["A2"]
    celda_meta.value = f"Generado el: {fecha_str} | Productos sugeridos: {len(items_a_pedir)}"
    celda_meta.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    celda_meta.alignment = align_center
    ws.row_dimensions[2].height = 20

    encabezados = [
        ("CÓDIGO", align_center),
        ("PRODUCTO", align_left),
        ("STOCK ACTUAL", align_center),
        ("STOCK MÍNIMO", align_center),
        ("STOCK MÁXIMO", align_center),
        ("CANTIDAD SUGERIDA A PEDIR", align_center),
        ("ESTADO", align_center),
    ]
    ws.row_dimensions[4].height = 24
    for col_idx, (enc, align) in enumerate(encabezados, start=1):
        cell = ws.cell(row=4, column=col_idx, value=enc)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align
        cell.border = borde_fino

    fila_actual = 5
    total_unidades = 0
    for item in items_a_pedir:
        cod = str(item["code"])
        nom = str(item["name"])
        stk = int(item["stock"])
        s_min = int(item["stock_minimo"])
        s_max = int(item["stock_maximo"])
        pedir = int(item["cantidad_sugerida_pedir"])
        estado = str(item["estado_inventario"])
        total_unidades += pedir

        ws.cell(row=fila_actual, column=1, value=cod).alignment = align_center
        ws.cell(row=fila_actual, column=2, value=nom).alignment = align_left
        ws.cell(row=fila_actual, column=3, value=stk).alignment = align_center
        ws.cell(row=fila_actual, column=4, value=s_min).alignment = align_center
        ws.cell(row=fila_actual, column=5, value=s_max).alignment = align_center

        c_pedir = ws.cell(row=fila_actual, column=6, value=pedir)
        c_pedir.alignment = align_center
        c_pedir.font = font_destacado

        c_est = ws.cell(row=fila_actual, column=7, value=estado)
        c_est.alignment = align_center
        c_est.font = font_agotado if estado == "AGOTADO" else font_quiebre

        for col_idx in range(1, 8):
            c = ws.cell(row=fila_actual, column=col_idx)
            if col_idx not in (6, 7):
                c.font = font_datos
            c.border = borde_fino

        ws.row_dimensions[fila_actual].height = 20
        fila_actual += 1

    if items_a_pedir:
        ws.row_dimensions[fila_actual].height = 24
        ws.cell(row=fila_actual, column=1, value="").border = borde_totales
        cell_lbl = ws.cell(row=fila_actual, column=2, value="TOTAL UNIDADES A PEDIR:")
        cell_lbl.font = Font(name="Segoe UI", size=10, bold=True)
        cell_lbl.alignment = align_right
        cell_lbl.border = borde_totales
        for col_idx in range(3, 6):
            ws.cell(row=fila_actual, column=col_idx, value="").border = borde_totales
        cell_tot = ws.cell(row=fila_actual, column=6, value=total_unidades)
        cell_tot.font = Font(name="Segoe UI", size=11, bold=True, color="0F766E")
        cell_tot.alignment = align_center
        cell_tot.border = borde_totales
        ws.cell(row=fila_actual, column=7, value="").border = borde_totales
    else:
        ws.merge_cells("A5:G5")
        c_vacio = ws["A5"]
        c_vacio.value = "No hay productos que requieran reposición actualmente."
        c_vacio.font = font_datos
        c_vacio.alignment = align_center
        ws.row_dimensions[5].height = 26

    anchos = {1: 14, 2: 45, 3: 16, 4: 16, 5: 16, 6: 28, 7: 20}
    for col_idx, ancho in anchos.items():
        letra = get_column_letter(col_idx)
        ws.column_dimensions[letra].width = ancho

    if ruta_salida is None:
        carpeta = Path(ruta_excel_base).parent / "Fabrica" / "Pedidos_Sugeridos"
        carpeta.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        ruta_final = carpeta / f"Pedido_Fabrica_{timestamp}.xlsx"
    else:
        ruta_final = Path(ruta_salida)
        ruta_final.parent.mkdir(parents=True, exist_ok=True)

    wb.save(str(ruta_final))
    wb.close()
    return ruta_final
