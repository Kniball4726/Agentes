"""Generación y exportación de informes en formatos Excel (.xls) y PDF (ReportLab)."""

from copy import copy
from pathlib import Path
from typing import Dict, Literal, Optional
from xml.sax.saxutils import escape

import openpyxl
import xlwt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_ALINEACIONES_PDF: dict[str, Literal[0, 1, 2]] = {
    "left": 0,
    "center": 1,
    "right": 2,
}


def copiar_estilo_fila(
    hoja_plantilla,
    hoja_salida,
    fila_plantilla: int,
    fila_salida: int,
    cantidad_columnas: int,
    columnas_plantilla: tuple[int, int, int, int],
) -> None:
    """Copia la altura y estilos de celda desde una fila de plantilla a la fila de salida."""
    hoja_salida.row_dimensions[fila_salida].height = (
        hoja_plantilla.row_dimensions[fila_plantilla].height
    )
    for columna in range(1, cantidad_columnas + 1):
        origen = hoja_plantilla.cell(
            fila_plantilla, columnas_plantilla[columna - 1]
        )
        destino = hoja_salida.cell(fila_salida, columna)
        if origen.has_style:
            destino.font = copy(origen.font)
            destino.fill = copy(origen.fill)
            destino.border = copy(origen.border)
            destino.alignment = copy(origen.alignment)
            destino.number_format = origen.number_format
            destino.protection = copy(origen.protection)


def guardar_informe_xls(libro: openpyxl.Workbook, ruta: Path) -> None:
    hoja_origen = libro.active
    if hoja_origen.max_row > 65_536:
        raise ValueError(
            "El informe supera el límite de 65.536 filas del formato XLS."
        )

    libro_xls = xlwt.Workbook(encoding="utf-8")
    hoja_xls = libro_xls.add_sheet(hoja_origen.title[:31])
    colores = {
        "17365D": 0x21,
        "1F4E78": 0x22,
        "D9EAF7": 0x23,
        "F3F7FB": 0x24,
        "D6DEE8": 0x25,
        "5B6573": 0x26,
        "E2F0D9": 0x27,
        "FCE4D6": 0x28,
        "9C0006": 0x29,
        "263445": 0x2A,
        "EAF0F6": 0x2B,
        "7F8995": 0x2C,
    }
    rgb_por_indice = {
        0x21: (23, 54, 93),
        0x22: (31, 78, 120),
        0x23: (217, 234, 247),
        0x24: (243, 247, 251),
        0x25: (214, 222, 232),
        0x26: (91, 101, 115),
        0x27: (226, 240, 217),
        0x28: (252, 228, 214),
        0x29: (156, 0, 6),
        0x2A: (38, 52, 69),
        0x2B: (234, 240, 246),
        0x2C: (127, 137, 149),
    }
    for indice, rgb in rgb_por_indice.items():
        libro_xls.set_colour_RGB(indice, *rgb)

    siguiente_color = 0x2D
    for fila in hoja_origen.iter_rows():
        for celda in fila:
            colores_celda = (
                celda.font.color,
                celda.fill.fgColor,
                celda.fill.bgColor,
                *(
                    getattr(getattr(celda.border, lado), "color", None)
                    for lado in ("left", "right", "top", "bottom")
                ),
            )
            for color in colores_celda:
                if color is None or color.type != "rgb" or not color.rgb:
                    continue
                codigo_rgb = str(color.rgb)[-6:].upper()
                if codigo_rgb in colores:
                    continue
                if siguiente_color > 0x3F:
                    break
                colores[codigo_rgb] = siguiente_color
                rgb = tuple(
                    int(codigo_rgb[indice:indice + 2], 16)
                    for indice in (0, 2, 4)
                )
                rgb_por_indice[siguiente_color] = rgb
                libro_xls.set_colour_RGB(siguiente_color, *rgb)
                siguiente_color += 1

    colores_borde = {
        "hair": xlwt.Borders.HAIR,
        "thin": xlwt.Borders.THIN,
        "medium": xlwt.Borders.MEDIUM,
        "dashed": xlwt.Borders.DASHED,
        "dotted": xlwt.Borders.DOTTED,
        "thick": xlwt.Borders.THICK,
    }
    horizontal = {
        "left": xlwt.Alignment.HORZ_LEFT,
        "center": xlwt.Alignment.HORZ_CENTER,
        "right": xlwt.Alignment.HORZ_RIGHT,
    }
    vertical = {
        "top": xlwt.Alignment.VERT_TOP,
        "center": xlwt.Alignment.VERT_CENTER,
        "bottom": xlwt.Alignment.VERT_BOTTOM,
    }
    estilos: Dict[int, xlwt.XFStyle] = {}

    def color_index(color) -> Optional[int]:
        if color is None or color.type != "rgb" or not color.rgb:
            return None
        return colores.get(str(color.rgb)[-6:].upper())

    def estilo_para(celda) -> xlwt.XFStyle:
        estilo = estilos.get(celda.style_id)
        if estilo is not None:
            return estilo

        estilo = xlwt.XFStyle()
        fuente = xlwt.Font()
        fuente.name = celda.font.name or "Arial"
        fuente.height = int((celda.font.sz or 10) * 20)
        fuente.bold = bool(celda.font.bold)
        fuente.italic = bool(celda.font.italic)
        fuente.struck_out = bool(celda.font.strike)
        if celda.font.underline:
            fuente.underline = xlwt.Font.UNDERLINE_SINGLE
        indice_color = color_index(celda.font.color)
        if indice_color is not None:
            fuente.colour_index = indice_color
        estilo.font = fuente

        alineacion = xlwt.Alignment()
        alineacion.horz = horizontal.get(celda.alignment.horizontal, xlwt.Alignment.HORZ_GENERAL)
        alineacion.vert = vertical.get(celda.alignment.vertical, xlwt.Alignment.VERT_BOTTOM)
        alineacion.wrap = bool(celda.alignment.wrap_text)
        alineacion.shrink_to_fit = bool(celda.alignment.shrink_to_fit)
        estilo.alignment = alineacion

        bordes = xlwt.Borders()
        for nombre in ("left", "right", "top", "bottom"):
            borde = getattr(celda.border, nombre)
            setattr(
                bordes,
                nombre,
                colores_borde.get(
                    getattr(borde, "style", None),
                    xlwt.Borders.NO_LINE,
                ),
            )
            indice_color = color_index(getattr(borde, "color", None))
            if indice_color is not None:
                setattr(bordes, f"{nombre}_colour", indice_color)
        estilo.borders = bordes

        patron = xlwt.Pattern()
        if celda.fill.fill_type == "solid":
            patron.pattern = xlwt.Pattern.SOLID_PATTERN
            indice_color = color_index(celda.fill.fgColor)
            rgb_relleno = rgb_por_indice.get(indice_color)
            luminancia = (
                sum(
                    peso * canal
                    for peso, canal in zip(
                        (0.2126, 0.7152, 0.0722),
                        rgb_relleno,
                    )
                )
                if rgb_relleno
                else 255
            )
            patron.pattern_fore_colour = (
                indice_color
                if indice_color is not None and luminancia >= 190
                else 9
            )
        estilo.pattern = patron
        estilo.num_format_str = celda.number_format or "General"
        estilos[celda.style_id] = estilo
        return estilo

    for columna in range(1, hoja_origen.max_column + 1):
        dimension = hoja_origen.column_dimensions[
            openpyxl.utils.get_column_letter(columna)
        ]
        ancho = dimension.width
        if ancho:
            hoja_xls.col(columna - 1).width = min(int(ancho * 256), 65_535)
        hoja_xls.col(columna - 1).hidden = dimension.hidden

    rangos_combinados = {}
    celdas_combinadas = set()
    for rango in hoja_origen.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rango.bounds
        rangos_combinados[(min_row, min_col)] = (
            min_row - 1,
            max_row - 1,
            min_col - 1,
            max_col - 1,
        )
        for fila in range(min_row, max_row + 1):
            for columna in range(min_col, max_col + 1):
                if (fila, columna) != (min_row, min_col):
                    celdas_combinadas.add((fila, columna))

    for fila in hoja_origen.iter_rows():
        for celda in fila:
            if (celda.row, celda.column) in celdas_combinadas:
                continue
            estilo = estilo_para(celda)
            valor = celda.value if celda.value is not None else ""
            rango = rangos_combinados.get((celda.row, celda.column))
            if rango:
                hoja_xls.write_merge(*rango, valor, estilo)
            else:
                hoja_xls.write(celda.row - 1, celda.column - 1, valor, estilo)

    for fila, dimension in hoja_origen.row_dimensions.items():
        if dimension.height is not None:
            hoja_xls.row(fila - 1).height = int(dimension.height * 20)

    hoja_xls.portrait = hoja_origen.page_setup.orientation == "portrait"
    hoja_xls.paper_size_code = int(hoja_origen.page_setup.paperSize)
    hoja_xls.left_margin = hoja_origen.page_margins.left
    hoja_xls.right_margin = hoja_origen.page_margins.right
    hoja_xls.top_margin = hoja_origen.page_margins.top
    hoja_xls.bottom_margin = hoja_origen.page_margins.bottom
    hoja_xls.fit_width_to_pages = hoja_origen.page_setup.fitToWidth
    hoja_xls.fit_height_to_pages = hoja_origen.page_setup.fitToHeight
    if hoja_origen.freeze_panes:
        fila_inmovilizada, columna_inmovilizada = (
            openpyxl.utils.cell.coordinate_to_tuple(hoja_origen.freeze_panes)
        )
        if fila_inmovilizada > 1 or columna_inmovilizada > 1:
            hoja_xls.set_panes_frozen(True)
            if fila_inmovilizada > 1:
                hoja_xls.set_horz_split_pos(fila_inmovilizada - 1)
            if columna_inmovilizada > 1:
                hoja_xls.set_vert_split_pos(columna_inmovilizada - 1)

    libro_xls.save(str(ruta))


def guardar_informe_pdf(libro: openpyxl.Workbook, ruta: Path) -> None:
    """Convierte el contenido visual del informe openpyxl a un documento PDF profesional."""
    hoja = libro.active
    if hoja is None:
        raise ValueError("El informe no contiene una hoja activa para exportar.")

    def color_excel(color) -> colors.Color | None:
        if color is None or color.type != "rgb" or not color.rgb:
            return None
        codigo = str(color.rgb)[-6:]
        try:
            return colors.HexColor(f"#{codigo}")
        except ValueError:
            return None

    def formatear_valor(celda) -> str:
        val = celda.value
        if val is None:
            return ""
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        res = escape(str(val))
        return res.replace("\n", "<br/>")

    ancho_util = A4[0] - 4 * cm  # 481.8898 pt

    fila_inicio_pedidos = None
    fila_inicio_no_descontados = None

    for num_f in range(1, hoja.max_row + 1):
        val_a = str(hoja.cell(num_f, 1).value or "").strip().upper()
        if "PEDIDOS CON PRODUCTOS DESCONTADOS" in val_a and fila_inicio_pedidos is None:
            fila_inicio_pedidos = num_f
        elif "PRODUCTOS NO DESCONTADOS" in val_a and fila_inicio_no_descontados is None:
            fila_inicio_no_descontados = num_f

    tiene_secciones = (fila_inicio_pedidos is not None) or (fila_inicio_no_descontados is not None)

    def pie_pagina(canvas, documento) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#5B6573"))
        canvas.setStrokeColor(colors.HexColor("#D6DEE8"))
        canvas.setLineWidth(0.4)
        canvas.line(2 * cm, 1.8 * cm + 6, A4[0] - 2 * cm, 1.8 * cm + 6)
        canvas.drawString(
            2 * cm,
            1.8 * cm - 6,
            "StockAgent · Informe de Control de Stock",
        )
        canvas.drawRightString(
            A4[0] - 2 * cm,
            1.8 * cm - 6,
            f"Página {documento.page}",
        )
        canvas.restoreState()

    elementos = []

    if tiene_secciones:
        fin_catalogo = (
            fila_inicio_pedidos
            if fila_inicio_pedidos is not None
            else (
                fila_inicio_no_descontados
                if fila_inicio_no_descontados is not None
                else hoja.max_row + 1
            )
        )

        ancho_codigo = 2.4 * cm
        ancho_descuento = 3.6 * cm
        ancho_producto = ancho_util - ancho_codigo - ancho_descuento
        anchos_catalogo = [ancho_codigo, ancho_producto, ancho_descuento]

        contenido_cat = []
        estilos_cat = [
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]

        idx_cat = 0
        for num_f in range(1, fin_catalogo):
            celdas_fila = [hoja.cell(num_f, c) for c in (1, 2, 3)]
            textos_fila = [formatear_valor(c) for c in celdas_fila]

            if all(not t for t in textos_fila):
                continue

            fila_elems = []
            for col_idx, c in enumerate(celdas_fila):
                align_str = c.alignment.horizontal or "left"
                align_code = _ALINEACIONES_PDF.get(align_str, 0)
                font_name = "Helvetica-Bold" if c.font and c.font.bold else "Helvetica"
                font_size = c.font.size if c.font and c.font.size else 9
                color_texto = color_excel(c.font.color) if c.font else colors.black
                if not color_texto:
                    color_texto = colors.black

                st_par = ParagraphStyle(
                    name=f"cat_{num_f}_{col_idx}",
                    fontName=font_name,
                    fontSize=font_size,
                    leading=font_size + 3,
                    textColor=color_texto,
                    alignment=align_code,
                )
                fila_elems.append(Paragraph(formatear_valor(c), st_par))

                bg = color_excel(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
                if bg:
                    estilos_cat.append(("BACKGROUND", (col_idx, idx_cat), (col_idx, idx_cat), bg))

                if c.border:
                    for side_name, side_obj in [
                        ("TOP", c.border.top),
                        ("BOTTOM", c.border.bottom),
                        ("LEFT", c.border.left),
                        ("RIGHT", c.border.right),
                    ]:
                        if side_obj and side_obj.style:
                            b_col = color_excel(side_obj.color) or colors.HexColor("#D6DEE8")
                            b_w = 1.0 if side_obj.style in ("medium", "thick") else 0.5
                            estilos_cat.append(
                                (f"{side_name}GRID" if side_name in ("TOP", "BOTTOM") else side_name,
                                 (col_idx, idx_cat), (col_idx, idx_cat), b_w, b_col)
                            )

            contenido_cat.append(fila_elems)
            idx_cat += 1

        for r_m in hoja.merged_cells.ranges:
            if r_m.min_row < fin_catalogo and r_m.max_row < fin_catalogo:
                min_c = r_m.min_col - 1
                max_c = min(r_m.max_col - 1, 2)
                r_orig_min = r_m.min_row
                r_orig_max = r_m.max_row
                r_target_min = None
                r_target_max = None
                cur_i = 0
                for nf in range(1, fin_catalogo):
                    row_empty = all(not formatear_valor(hoja.cell(nf, c)) for c in (1, 2, 3))
                    if not row_empty:
                        if nf == r_orig_min:
                            r_target_min = cur_i
                        if nf == r_orig_max:
                            r_target_max = cur_i
                        cur_i += 1
                if r_target_min is not None and r_target_max is not None:
                    estilos_cat.append(("SPAN", (min_c, r_target_min), (max_c, r_target_max)))

        if contenido_cat:
            t_cat = Table(
                contenido_cat,
                colWidths=anchos_catalogo,
                repeatRows=min(3, len(contenido_cat)),
                hAlign="LEFT",
                splitByRow=1,
            )
            t_cat.setStyle(TableStyle(estilos_cat))
            elementos.append(t_cat)

        if fila_inicio_pedidos is not None:
            elementos.append(Spacer(1, 14))
            fin_pedidos = (
                fila_inicio_no_descontados
                if fila_inicio_no_descontados is not None
                else hoja.max_row + 1
            )

            ancho_col1 = 2.4 * cm
            ancho_col2 = ancho_util - ancho_col1
            anchos_pedidos = [ancho_col1, ancho_col2]

            contenido_ped = []
            estilos_ped = [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]

            idx_ped = 0
            for num_f in range(fila_inicio_pedidos, fin_pedidos):
                c1 = hoja.cell(num_f, 1)
                c2 = hoja.cell(num_f, 2)
                t1 = formatear_valor(c1)
                t2 = formatear_valor(c2)

                if not t1 and not t2:
                    continue

                fila_elems = []
                for col_idx, c in enumerate([c1, c2]):
                    align_str = c.alignment.horizontal or "left"
                    align_code = _ALINEACIONES_PDF.get(align_str, 0)
                    font_name = "Helvetica-Bold" if c.font and c.font.bold else "Helvetica"
                    font_size = c.font.size if c.font and c.font.size else 9
                    color_texto = color_excel(c.font.color) if c.font else colors.black
                    if not color_texto:
                        color_texto = colors.black

                    st_par = ParagraphStyle(
                        name=f"ped_{num_f}_{col_idx}",
                        fontName=font_name,
                        fontSize=font_size,
                        leading=font_size + 3,
                        textColor=color_texto,
                        alignment=align_code,
                    )
                    fila_elems.append(Paragraph(formatear_valor(c), st_par))

                    bg = color_excel(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
                    if bg:
                        estilos_ped.append(("BACKGROUND", (col_idx, idx_ped), (col_idx, idx_ped), bg))

                    if c.border:
                        for side_name, side_obj in [
                            ("TOP", c.border.top),
                            ("BOTTOM", c.border.bottom),
                            ("LEFT", c.border.left),
                            ("RIGHT", c.border.right),
                        ]:
                            if side_obj and side_obj.style:
                                b_col = color_excel(side_obj.color) or colors.HexColor("#D6DEE8")
                                b_w = 1.0 if side_obj.style in ("medium", "thick") else 0.5
                                estilos_ped.append(
                                    (f"{side_name}GRID" if side_name in ("TOP", "BOTTOM") else side_name,
                                     (col_idx, idx_ped), (col_idx, idx_ped), b_w, b_col)
                                )

                contenido_ped.append(fila_elems)
                idx_ped += 1

            for r_m in hoja.merged_cells.ranges:
                if fila_inicio_pedidos <= r_m.min_row < fin_pedidos:
                    min_c = r_m.min_col - 1
                    max_c = min(r_m.max_col - 1, 1)
                    r_target_min = None
                    r_target_max = None
                    cur_i = 0
                    for nf in range(fila_inicio_pedidos, fin_pedidos):
                        row_empty = not formatear_valor(hoja.cell(nf, 1)) and not formatear_valor(hoja.cell(nf, 2))
                        if not row_empty:
                            if nf == r_m.min_row:
                                r_target_min = cur_i
                            if nf == r_m.max_row:
                                r_target_max = cur_i
                            cur_i += 1
                    if r_target_min is not None and r_target_max is not None:
                        estilos_ped.append(("SPAN", (min_c, r_target_min), (max_c, r_target_max)))

            if contenido_ped:
                t_ped = Table(
                    contenido_ped,
                    colWidths=anchos_pedidos,
                    repeatRows=min(2, len(contenido_ped)),
                    hAlign="LEFT",
                    splitByRow=1,
                )
                t_ped.setStyle(TableStyle(estilos_ped))
                elementos.append(t_ped)

        if fila_inicio_no_descontados is not None:
            elementos.append(Spacer(1, 14))
            ancho_col1 = 2.4 * cm
            ancho_col2 = ancho_util - ancho_col1
            anchos_nd = [ancho_col1, ancho_col2]

            contenido_nd = []
            estilos_nd = [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]

            idx_nd = 0
            for num_f in range(fila_inicio_no_descontados, hoja.max_row + 1):
                c1 = hoja.cell(num_f, 1)
                c2 = hoja.cell(num_f, 2)
                t1 = formatear_valor(c1)
                t2 = formatear_valor(c2)

                if not t1 and not t2:
                    continue

                fila_elems = []
                for col_idx, c in enumerate([c1, c2]):
                    align_str = c.alignment.horizontal or "left"
                    align_code = _ALINEACIONES_PDF.get(align_str, 0)
                    font_name = "Helvetica-Bold" if c.font and c.font.bold else "Helvetica"
                    font_size = c.font.size if c.font and c.font.size else 9
                    color_texto = color_excel(c.font.color) if c.font else colors.black
                    if not color_texto:
                        color_texto = colors.black

                    st_par = ParagraphStyle(
                        name=f"nd_{num_f}_{col_idx}",
                        fontName=font_name,
                        fontSize=font_size,
                        leading=font_size + 3,
                        textColor=color_texto,
                        alignment=align_code,
                    )
                    fila_elems.append(Paragraph(formatear_valor(c), st_par))

                    bg = color_excel(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
                    if bg:
                        estilos_nd.append(("BACKGROUND", (col_idx, idx_nd), (col_idx, idx_nd), bg))

                    if c.border:
                        for side_name, side_obj in [
                            ("TOP", c.border.top),
                            ("BOTTOM", c.border.bottom),
                            ("LEFT", c.border.left),
                            ("RIGHT", c.border.right),
                        ]:
                            if side_obj and side_obj.style:
                                b_col = color_excel(side_obj.color) or colors.HexColor("#D6DEE8")
                                b_w = 1.0 if side_obj.style in ("medium", "thick") else 0.5
                                estilos_nd.append(
                                    (f"{side_name}GRID" if side_name in ("TOP", "BOTTOM") else side_name,
                                     (col_idx, idx_nd), (col_idx, idx_nd), b_w, b_col)
                                )

                contenido_nd.append(fila_elems)
                idx_nd += 1

            for r_m in hoja.merged_cells.ranges:
                if r_m.min_row >= fila_inicio_no_descontados:
                    min_c = r_m.min_col - 1
                    max_c = min(r_m.max_col - 1, 1)
                    r_target_min = None
                    r_target_max = None
                    cur_i = 0
                    for nf in range(fila_inicio_no_descontados, hoja.max_row + 1):
                        row_empty = not formatear_valor(hoja.cell(nf, 1)) and not formatear_valor(hoja.cell(nf, 2))
                        if not row_empty:
                            if nf == r_m.min_row:
                                r_target_min = cur_i
                            if nf == r_m.max_row:
                                r_target_max = cur_i
                            cur_i += 1
                    if r_target_min is not None and r_target_max is not None:
                        estilos_nd.append(("SPAN", (min_c, r_target_min), (max_c, r_target_max)))

            if contenido_nd:
                t_nd = Table(
                    contenido_nd,
                    colWidths=anchos_nd,
                    repeatRows=min(2, len(contenido_nd)),
                    hAlign="LEFT",
                    splitByRow=1,
                )
                t_nd.setStyle(TableStyle(estilos_nd))
                elementos.append(t_nd)

    else:
        # Modo genérico (fallback)
        total_columnas = min(hoja.max_column, 6)
        anchos = [ancho_util / total_columnas] * total_columnas
        contenido = []
        estilos = [
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]

        for numero_fila, fila in enumerate(
            hoja.iter_rows(max_col=total_columnas), start=0
        ):
            fila_elementos = []
            for numero_columna, celda in enumerate(fila):
                alineacion_str = celda.alignment.horizontal or "left"
                alineacion_codigo = _ALINEACIONES_PDF.get(alineacion_str, 0)
                nombre_fuente = "Helvetica-Bold" if celda.font and celda.font.bold else "Helvetica"
                tamano_fuente = celda.font.size if celda.font and celda.font.size else 9
                color_texto = color_excel(celda.font.color) if celda.font else colors.black
                if not color_texto:
                    color_texto = colors.black

                estilo_parrafo = ParagraphStyle(
                    name=f"p_{numero_fila}_{numero_columna}",
                    fontName=nombre_fuente,
                    fontSize=tamano_fuente,
                    leading=tamano_fuente + 3,
                    textColor=color_texto,
                    alignment=alineacion_codigo,
                )
                fila_elementos.append(
                    Paragraph(formatear_valor(celda), estilo_parrafo)
                )

                color_fondo = color_excel(celda.fill.fgColor) if celda.fill and celda.fill.fill_type == "solid" else None
                if color_fondo:
                    estilos.append(
                        ("BACKGROUND", (numero_columna, numero_fila), (numero_columna, numero_fila), color_fondo)
                    )

            contenido.append(fila_elementos)

        tabla = Table(
            contenido,
            colWidths=anchos,
            repeatRows=min(3, len(contenido)),
            hAlign="LEFT",
            splitByRow=1,
        )
        tabla.setStyle(TableStyle(estilos))
        elementos.append(tabla)

    documento = SimpleDocTemplate(
        str(ruta),
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=ruta.stem,
    )
    documento.build(elementos, onFirstPage=pie_pagina, onLaterPages=pie_pagina)
