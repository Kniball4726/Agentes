"""Comparativa de stock digital vs. relevamiento físico y generación de reportes unificados."""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..remito_parser import normalizar_codigo


def comparar_stock_fisico_dict(
    catalogo: Dict[str, tuple],
    conteo_fisico: Dict[str, int],
) -> List[Dict]:
    """Compara un diccionario {codigo: stock_fisico} con el catálogo digital.

    Retorna una lista de diccionarios con:
        code, name, stock_digital, stock_fisico, diferencia, estado, en_catalogo
    """
    resultado: List[Dict] = []
    for codigo_raw, cant_fisica in conteo_fisico.items():
        cod = normalizar_codigo(codigo_raw)
        stk_fis = int(cant_fisica)
        if cod in catalogo:
            _, nom, stk_dig = catalogo[cod]
            stk_dig = int(stk_dig)
            dif = stk_fis - stk_dig
            if dif == 0:
                estado = "COINCIDE"
            elif dif < 0:
                estado = "FALTANTE"
            else:
                estado = "SOBRANTE"
            en_catalogo = True
        else:
            nom = "No registrado en catálogo"
            stk_dig = 0
            dif = stk_fis
            estado = "NO_CATALOGADO"
            en_catalogo = False

        resultado.append(
            {
                "code": cod,
                "name": nom,
                "stock_digital": stk_dig,
                "stock_fisico": stk_fis,
                "diferencia": dif,
                "estado": estado,
                "en_catalogo": en_catalogo,
            }
        )
    return resultado


def generar_documento_comparativo_excel(
    comparativa: List[Dict],
    ruta_excel_base: Path,
    product_store: Any = None,
    usuario: str = "operador",
    nombre_archivo_origen: str = "",
    ruta_salida: Path | str | None = None,
) -> Path:
    """Genera un archivo Excel unificado con la comparativa entre Stock Digital y Stock Físico."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Comparativa y Diferencias"

    fill_titulo = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_titulo = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_datos = Font(name="Segoe UI", size=10)
    font_coincide = Font(name="Segoe UI", size=10, bold=True, color="15803D")
    font_faltante = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    font_sobrante = Font(name="Segoe UI", size=10, bold=True, color="D97706")

    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    borde_totales = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A1:H1")
    celda_titulo = ws["A1"]
    celda_titulo.value = "COMPARATIVA UNIFICADA: STOCK DIGITAL VS. STOCK FÍSICO"
    celda_titulo.fill = fill_titulo
    celda_titulo.font = font_titulo
    celda_titulo.alignment = align_center
    ws.row_dimensions[1].height = 36

    fecha_str = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    con_dif = sum(1 for i in comparativa if i["diferencia"] != 0)
    ws.merge_cells("A2:H2")
    celda_meta = ws["A2"]
    celda_meta.value = (
        f"Fecha: {fecha_str} | Operario: {usuario} | Origen: {nombre_archivo_origen} | "
        f"Relevados: {len(comparativa)} | Con diferencias: {con_dif}"
    )
    celda_meta.font = Font(name="Segoe UI", size=9, italic=True, color="475569")
    celda_meta.alignment = align_center
    ws.row_dimensions[2].height = 20

    encabezados = [
        ("CÓDIGO", align_center),
        ("PRODUCTO", align_left),
        ("STOCK DIGITAL", align_center),
        ("STOCK FÍSICO", align_center),
        ("DIFERENCIA", align_center),
        ("ESTADO", align_center),
        ("OPERARIO", align_center),
        ("FECHA", align_center),
    ]
    ws.row_dimensions[4].height = 24
    for col_idx, (enc, align) in enumerate(encabezados, start=1):
        cell = ws.cell(row=4, column=col_idx, value=enc)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align
        cell.border = borde_fino

    fila_actual = 5
    for item in comparativa:
        cod = str(item["code"])
        nom = str(item["name"])
        stk_dig = int(item["stock_digital"])
        stk_fis = int(item["stock_fisico"])
        dif = int(item["diferencia"])
        estado = str(item["estado"])
        dif_str = f"+{dif}" if dif > 0 else str(dif)

        ws.cell(row=fila_actual, column=1, value=cod).alignment = align_center
        ws.cell(row=fila_actual, column=2, value=nom).alignment = align_left
        ws.cell(row=fila_actual, column=3, value=stk_dig).alignment = align_center
        ws.cell(row=fila_actual, column=4, value=stk_fis).alignment = align_center

        c_dif = ws.cell(row=fila_actual, column=5, value=dif_str)
        c_dif.alignment = align_center

        c_est = ws.cell(row=fila_actual, column=6, value=estado)
        c_est.alignment = align_center

        ws.cell(row=fila_actual, column=7, value=usuario).alignment = align_center
        ws.cell(row=fila_actual, column=8, value=fecha_str).alignment = align_center

        font_estado = (
            font_coincide
            if dif == 0
            else (font_faltante if dif < 0 else font_sobrante)
        )
        c_dif.font = font_estado
        c_est.font = font_estado

        for col_idx in range(1, 9):
            c = ws.cell(row=fila_actual, column=col_idx)
            if col_idx not in (5, 6):
                c.font = font_datos
            c.border = borde_fino

        ws.row_dimensions[fila_actual].height = 20
        fila_actual += 1

    if comparativa:
        ws.row_dimensions[fila_actual].height = 24
        for col_idx in range(1, 9):
            ws.cell(row=fila_actual, column=col_idx).border = borde_totales
        c_res = ws.cell(
            row=fila_actual,
            column=2,
            value=f"RESUMEN: {len(comparativa)} relevados ({con_dif} con diferencia)",
        )
        c_res.font = Font(name="Segoe UI", size=10, bold=True)

    anchos = {1: 14, 2: 44, 3: 16, 4: 16, 5: 16, 6: 18, 7: 18, 8: 22}
    for col_idx, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    # Segunda hoja: Catálogo Unificado Completo (Digital + Físico)
    if product_store is not None:
        ws_unif = wb.create_sheet(title="Catálogo Unificado")
        ws_unif.merge_cells("A1:F1")
        c_tit2 = ws_unif["A1"]
        c_tit2.value = "INVENTARIO UNIFICADO (CATÁLOGO COMPLETO)"
        c_tit2.fill = fill_titulo
        c_tit2.font = font_titulo
        c_tit2.alignment = align_center
        ws_unif.row_dimensions[1].height = 32

        enc_unif = [
            ("CÓDIGO", align_center),
            ("PRODUCTO", align_left),
            ("STOCK DIGITAL", align_center),
            ("STOCK FÍSICO", align_center),
            ("DIFERENCIA", align_center),
            ("ESTADO", align_center),
        ]
        ws_unif.row_dimensions[3].height = 22
        for col_idx, (enc, align) in enumerate(enc_unif, start=1):
            c = ws_unif.cell(row=3, column=col_idx, value=enc)
            c.fill = fill_header
            c.font = font_header
            c.alignment = align
            c.border = borde_fino

        mapa_relevados = {str(i["code"]): i for i in comparativa}
        productos_db = product_store.list_products()
        r_idx = 4
        for p in productos_db:
            if not p["active"]:
                continue
            cod = str(p["code"])
            nom = str(p["name"])
            stk_dig = int(p["stock"])
            if cod in mapa_relevados:
                stk_fis_val = int(mapa_relevados[cod]["stock_fisico"])
            elif p["stock_fisico"] is not None:
                stk_fis_val = int(p["stock_fisico"])
            else:
                stk_fis_val = None

            ws_unif.cell(row=r_idx, column=1, value=cod).alignment = align_center
            ws_unif.cell(row=r_idx, column=2, value=nom).alignment = align_left
            ws_unif.cell(row=r_idx, column=3, value=stk_dig).alignment = align_center

            if stk_fis_val is not None:
                dif = stk_fis_val - stk_dig
                est = "COINCIDE" if dif == 0 else ("FALTANTE" if dif < 0 else "SOBRANTE")
                ws_unif.cell(row=r_idx, column=4, value=stk_fis_val).alignment = align_center
                ws_unif.cell(row=r_idx, column=5, value=dif).alignment = align_center
                ws_unif.cell(row=r_idx, column=6, value=est).alignment = align_center
            else:
                ws_unif.cell(row=r_idx, column=4, value="-").alignment = align_center
                ws_unif.cell(row=r_idx, column=5, value="-").alignment = align_center
                ws_unif.cell(row=r_idx, column=6, value="SIN RELEVAR").alignment = align_center

            for col_idx in range(1, 7):
                cell_u = ws_unif.cell(row=r_idx, column=col_idx)
                cell_u.font = font_datos
                cell_u.border = borde_fino
            r_idx += 1

        for col_idx, ancho in {1: 14, 2: 44, 3: 16, 4: 16, 5: 16, 6: 18}.items():
            ws_unif.column_dimensions[get_column_letter(col_idx)].width = ancho

    if ruta_salida is None:
        carpeta = Path(ruta_excel_base).parent / "StockFisico" / "Comparativas"
        carpeta.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        ruta_final = carpeta / f"Comparativa_Stock_Fisico_{timestamp}.xlsx"
    else:
        ruta_final = Path(ruta_salida)
        ruta_final.parent.mkdir(parents=True, exist_ok=True)

    wb.save(str(ruta_final))
    wb.close()
    return ruta_final
