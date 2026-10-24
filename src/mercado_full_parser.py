"""
Módulo para parsear planillas Excel (.xlsx, .xls) de Mercado Envíos Full.
Extrae los códigos/SKU de los productos y las cantidades declaradas a enviar a Full.
"""

from pathlib import Path
from typing import Dict, List, Tuple
import openpyxl
import xlrd
from .remito_parser import normalizar_codigo


def parsear_archivo_mercado_full(ruta_archivo: str | Path) -> Tuple[Dict[str, int], List[Tuple[str, str, int]]]:
    """
    Parsea una planilla Excel (.xlsx o .xls) con columnas de Código/SKU y Cantidad de Mercado Full.
    
    Retorna:
        - items_consolidados: Dict[codigo, cantidad_total] acumulado por producto.
        - detalle: List[(codigo, descripcion, cantidad)] con el detalle fila a fila.
    """
    path = Path(ruta_archivo)
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")

    ext = path.suffix.lower()
    if ext not in (".xlsx", ".xls"):
        raise ValueError(f"Formato no soportado para Mercado Full: {ext}. Debe ser .xlsx o .xls")

    items_crudos: List[Tuple[str, str, int]] = []

    if ext == ".xls":
        libro = xlrd.open_workbook(str(path))
        hoja = libro.sheet_by_index(0)

        col_cod = -1
        col_desc = -1
        col_cant = -1
        fila_inicio = 0

        # Buscar encabezados en las primeras 20 filas
        for r in range(min(20, hoja.nrows)):
            fila_vals = [str(hoja.cell_value(r, c)).strip().upper() for c in range(hoja.ncols)]
            for c, val in enumerate(fila_vals):
                if any(k in val for k in ["SKU", "CODIGO", "CÓDIGO", "COD", "PUBLICACION", "PUBLICACIÓN", "ARTICULO", "ARTÍCULO", "MLC", "MLA"]) and col_cod == -1:
                    col_cod = c
                elif any(k in val for k in ["TITULO", "TÍTULO", "DESCRIPCION", "DESCRIPCIÓN", "PRODUCTO", "NOMBRE", "DETALLE"]) and col_desc == -1:
                    col_desc = c
                elif any(k in val for k in ["CANTIDAD", "CANT", "UNIDADES", "ENVIADAS", "ENVIAR", "DECLARADA", "DECLARADAS", "PIEZAS", "BULTOS"]) and col_cant == -1:
                    col_cant = c
            if col_cod != -1 and col_cant != -1:
                fila_inicio = r + 1
                break

        # Fallback si no hay encabezados explícitos reconocidos
        if col_cod == -1 or col_cant == -1:
            col_cod = 0
            col_desc = 1 if hoja.ncols > 2 else -1
            col_cant = 2 if hoja.ncols > 2 else (1 if hoja.ncols > 1 else 0)
            fila_inicio = 0

        for r in range(fila_inicio, hoja.nrows):
            val_cod = hoja.cell_value(r, col_cod)
            val_cant = hoja.cell_value(r, col_cant)
            val_desc = hoja.cell_value(r, col_desc) if col_desc != -1 else ""

            if val_cod is None or str(val_cod).strip() == "":
                continue

            str_cod = str(val_cod).strip()
            # Ignorar filas de totales o encabezados residuales
            if str_cod.upper() in ("TOTAL", "TOTALES", "SUBTOTAL", "RESUMEN"):
                continue

            cod = normalizar_codigo(str_cod)
            try:
                cant = int(float(val_cant))
            except (ValueError, TypeError):
                continue

            if cant > 0:
                items_crudos.append((cod, str(val_desc).strip(), cant))

    else:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        hoja = wb.active
        if hoja is None:
            return {}, []

        col_cod = -1
        col_desc = -1
        col_cant = -1
        fila_inicio = 1

        for r in range(1, min(20, hoja.max_row + 1)):
            fila_vals = [str(hoja.cell(r, c).value or "").strip().upper() for c in range(1, hoja.max_column + 1)]
            for idx, val in enumerate(fila_vals):
                c = idx + 1
                if any(k in val for k in ["SKU", "CODIGO", "CÓDIGO", "COD", "PUBLICACION", "PUBLICACIÓN", "ARTICULO", "ARTÍCULO", "MLC", "MLA"]) and col_cod == -1:
                    col_cod = c
                elif any(k in val for k in ["TITULO", "TÍTULO", "DESCRIPCION", "DESCRIPCIÓN", "PRODUCTO", "NOMBRE", "DETALLE"]) and col_desc == -1:
                    col_desc = c
                elif any(k in val for k in ["CANTIDAD", "CANT", "UNIDADES", "ENVIADAS", "ENVIAR", "DECLARADA", "DECLARADAS", "PIEZAS", "BULTOS"]) and col_cant == -1:
                    col_cant = c
            if col_cod != -1 and col_cant != -1:
                fila_inicio = r + 1
                break

        if col_cod == -1 or col_cant == -1:
            col_cod = 1
            col_desc = 2 if hoja.max_column > 2 else -1
            col_cant = 3 if hoja.max_column > 2 else (2 if hoja.max_column > 1 else 1)
            fila_inicio = 1

        for r in range(fila_inicio, hoja.max_row + 1):
            val_cod = hoja.cell(r, col_cod).value
            val_cant = hoja.cell(r, col_cant).value
            val_desc = hoja.cell(r, col_desc).value if col_desc != -1 else ""

            if val_cod is None or str(val_cod).strip() == "":
                continue

            str_cod = str(val_cod).strip()
            if str_cod.upper() in ("TOTAL", "TOTALES", "SUBTOTAL", "RESUMEN"):
                continue

            cod = normalizar_codigo(str_cod)
            try:
                cant = int(float(val_cant))
            except (ValueError, TypeError):
                continue

            if cant > 0:
                items_crudos.append((cod, str(val_desc or "").strip(), cant))

    # Consolidar por código sumando cantidades
    items_consolidados: Dict[str, int] = {}
    for cod, _, cant in items_crudos:
        items_consolidados[cod] = items_consolidados.get(cod, 0) + cant

    return items_consolidados, items_crudos
