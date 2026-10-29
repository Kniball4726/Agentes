"""Carga el catálogo desde SQLite y genera informes Excel con la plantilla visual."""

from dataclasses import dataclass
from copy import copy
from datetime import datetime
import os
from pathlib import Path
import shutil
import sqlite3
import tempfile
from typing import Dict, List, Literal, Optional, Tuple
import unicodedata
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .consolidator import ItemConsolidado
from .product_store import ProductStore
from .remito_parser import ResultadoRemito, normalizar_codigo
from .stock.models import RegistroDescuento, ReporteStock
from .stock.reports import copiar_estilo_fila, guardar_informe_pdf, guardar_informe_xls
from .stock.replenishment import generar_pedido_sugerido_fabrica_excel
from .stock.physical_inventory import (
    comparar_stock_fisico_dict,
    generar_documento_comparativo_excel,
)


class StockManager:
    def __init__(
        self,
        ruta_excel_base: str | Path,
        nombre_hoja: str = "Stock",
        ruta_base_datos: str | Path | None = None,
    ):
        self.ruta_excel_base = Path(ruta_excel_base)
        self.nombre_hoja = nombre_hoja
        carpeta_base_datos = (
            self.ruta_excel_base.parent.parent
            if self.ruta_excel_base.parent.name.casefold() == "descuentos"
            else self.ruta_excel_base.parent
        )
        self.ruta_base_datos = (
            Path(ruta_base_datos)
            if ruta_base_datos is not None
            else carpeta_base_datos / "data" / "productos.sqlite3"
        )
        self.product_store = ProductStore(self.ruta_base_datos)
        # Lista ordenada de productos únicos: (fila_original, codigo, nombre, stock_actual)
        self.productos_base: List[Tuple[int, str, str, int]] = []
        self.filas_catalogo: List[Tuple[int, str, str, int]] = []
        self.catalogo: Dict[str, Tuple[int, str, int]] = {}
        self._filas_por_codigo: Dict[str, List[int]] = {}
        self._col_codigo = 1
        self._col_nombre = 2
        self._col_stock = 3
        self._col_stock_inicial: Optional[int] = None
        self._col_descuento: Optional[int] = None
        self._col_stock_final: Optional[int] = None
        self._fila_inicio_datos = 3
        self._cargar_catalogo()

    @staticmethod
    def _normalizar_encabezado(valor) -> str:
        texto = unicodedata.normalize("NFKD", str(valor or ""))
        return "".join(caracter for caracter in texto if not unicodedata.combining(caracter)).strip().upper()

    def _detectar_esquema(self, sheet) -> None:
        encabezados_codigo = {"CODIGO", "COD"}
        encabezados_nombre = {"PRODUCTO", "PRODUCTOS", "DESCRIPCION", "ARTICULO"}
        fila_encabezados = None
        columnas = {}
        filas_stock = []

        for fila in range(1, min(sheet.max_row, 14) + 1):
            fila_actual = {}
            for columna in range(1, sheet.max_column + 1):
                valor = self._normalizar_encabezado(sheet.cell(fila, columna).value)
                if valor in encabezados_codigo:
                    fila_actual["codigo"] = columna
                elif valor in encabezados_nombre:
                    fila_actual["nombre"] = columna
                elif valor == "STOCK INICIAL":
                    self._col_stock_inicial = columna
                    filas_stock.append(fila)
                elif valor == "DESCUENTO":
                    self._col_descuento = columna
                    filas_stock.append(fila)
                elif valor == "STOCK FINAL":
                    self._col_stock_final = columna
                    filas_stock.append(fila)

            if "codigo" in fila_actual and "nombre" in fila_actual:
                fila_encabezados = fila
                columnas = fila_actual

        if fila_encabezados is None:
            self._col_codigo = 1
            self._col_nombre = 2
            self._col_stock = 3
            self._col_stock_inicial = None
            self._col_descuento = None
            self._col_stock_final = None
            self._fila_inicio_datos = 3
            return

        self._col_codigo = columnas["codigo"]
        self._col_nombre = columnas["nombre"]
        self._col_stock = self._col_stock_inicial or 3
        self._fila_inicio_datos = max(3, fila_encabezados + 1, *(fila + 1 for fila in filas_stock))

    def _leer_stock(self, sheet, fila: int) -> int:
        columnas_stock = [self._col_stock_final, self._col_stock_inicial or self._col_stock]
        for columna in columnas_stock:
            if columna is None:
                continue
            valor = sheet.cell(row=fila, column=columna).value
            if valor is None or not str(valor).strip():
                continue
            try:
                return int(round(float(str(valor).replace(",", "."))))
            except (ValueError, TypeError):
                continue
        return 0

    def _cargar_catalogo(self):
        """Importa el catálogo inicial y después carga productos desde SQLite."""
        if not self.ruta_excel_base.exists():
            raise FileNotFoundError(f"No se encontró el archivo de stock en: {self.ruta_excel_base}")

        self.productos_base.clear()
        self.filas_catalogo.clear()
        self.catalogo.clear()
        self._filas_por_codigo.clear()

        wb = openpyxl.load_workbook(str(self.ruta_excel_base), data_only=True)
        sheet = wb[self.nombre_hoja] if self.nombre_hoja in wb.sheetnames else wb.active
        if sheet is None:
            wb.close()
            raise ValueError(
                f"No se encontró una hoja activa en {self.ruta_excel_base}."
            )
        self._detectar_esquema(sheet)
        productos_plantilla = []
        for fila in range(self._fila_inicio_datos, sheet.max_row + 1):
            val_cod = sheet.cell(row=fila, column=self._col_codigo).value
            if val_cod is not None and str(val_cod).strip() != "":
                cod = normalizar_codigo(str(val_cod))
                nombre = str(sheet.cell(row=fila, column=self._col_nombre).value or "").strip()
                if not nombre:
                    continue
                stock_num = self._leer_stock(sheet, fila)
                productos_plantilla.append((fila, cod, nombre, stock_num))
        wb.close()

        self.product_store.import_template_once(productos_plantilla)
        primera_aparicion: set[str] = set()
        for indice, producto in enumerate(self.product_store.list_catalog_rows()):
            cod = str(producto["code"])
            nombre = str(producto["display_name"])
            nombre_catalogo = str(producto["name"])
            stock_num = int(producto["stock"])
            fila_plantilla = producto["template_row"]
            fila_estilo = int(fila_plantilla) if fila_plantilla is not None else 0
            fila_catalogo = indice + self._fila_inicio_datos

            self.filas_catalogo.append(
                (fila_estilo, cod, nombre, stock_num)
            )
            self._filas_por_codigo.setdefault(cod, []).append(fila_catalogo)
            if cod in primera_aparicion:
                continue

            primera_aparicion.add(cod)
            self.productos_base.append(
                (fila_catalogo, cod, nombre_catalogo, stock_num)
            )
            self.catalogo[cod] = (fila_catalogo, nombre_catalogo, stock_num)

    def listar_productos(self) -> List[sqlite3.Row]:
        """Devuelve todos los productos para la pantalla de administración."""
        return self.product_store.list_products()

    def actualizar_producto(
        self,
        codigo_actual: str,
        codigo_nuevo: str,
        nombre: str,
        stock: int,
        stock_minimo: Optional[int] = None,
        stock_maximo: Optional[int] = None,
    ) -> None:
        codigo_actual = normalizar_codigo(codigo_actual)
        codigo_nuevo = normalizar_codigo(codigo_nuevo)
        nombre = nombre.strip()
        if not codigo_actual or not codigo_nuevo or not nombre:
            raise ValueError("El código y el nombre del producto son obligatorios.")

        contenido_original = self.ruta_excel_base.read_bytes()
        respaldo = self.ruta_excel_base.with_name(
            f"{self.ruta_excel_base.stem}_backup{self.ruta_excel_base.suffix}"
        )
        plantilla_actualizada = False

        def sincronizar_plantilla() -> list[int]:
            nonlocal plantilla_actualizada
            shutil.copy2(self.ruta_excel_base, respaldo)
            libro = openpyxl.load_workbook(str(self.ruta_excel_base))
            try:
                hoja = (
                    libro[self.nombre_hoja]
                    if self.nombre_hoja in libro.sheetnames
                    else libro.active
                )
                if hoja is None:
                    raise ValueError("No se encontró una hoja para actualizar.")

                filas_producto = [
                    fila
                    for fila in range(self._fila_inicio_datos, hoja.max_row + 1)
                    if normalizar_codigo(
                        str(hoja.cell(fila, self._col_codigo).value or "")
                    )
                    == codigo_actual
                ]
                filas_catalogo_producto = [
                    producto
                    for producto in self.product_store.list_catalog_rows(
                        include_inactive=True
                    )
                    if str(producto["code"]) == codigo_actual
                ]
                if not filas_producto or len(filas_producto) != len(
                    filas_catalogo_producto
                ):
                    raise RuntimeError(
                        f"Las filas de '{codigo_actual}' en Productos.xlsx "
                        "no coinciden con el catálogo."
                    )

                codigo_excel = (
                    int(codigo_nuevo)
                    if codigo_nuevo.isdigit()
                    else codigo_nuevo
                )
                for fila in filas_producto:
                    hoja.cell(fila, self._col_codigo, codigo_excel)
                    hoja.cell(fila, self._col_nombre, nombre)
                    if self._col_stock_final is not None:
                        columna_inicial = self._col_stock_inicial or self._col_stock
                        hoja.cell(fila, columna_inicial, stock)
                        if self._col_descuento is not None:
                            hoja.cell(fila, self._col_descuento, None)
                        hoja.cell(fila, self._col_stock_final, stock)
                    else:
                        hoja.cell(fila, self._col_stock, stock)

                with tempfile.NamedTemporaryFile(
                    suffix=self.ruta_excel_base.suffix,
                    dir=self.ruta_excel_base.parent,
                    delete=False,
                ) as temporal:
                    ruta_temporal = Path(temporal.name)
                try:
                    libro.save(ruta_temporal)
                finally:
                    libro.close()
                try:
                    os.replace(ruta_temporal, self.ruta_excel_base)
                    plantilla_actualizada = True
                finally:
                    if ruta_temporal.exists():
                        ruta_temporal.unlink()

            except Exception:
                libro.close()
                raise

            libro_verificacion = openpyxl.load_workbook(
                str(self.ruta_excel_base), data_only=True, read_only=True
            )
            try:
                hoja_verificacion = (
                    libro_verificacion[self.nombre_hoja]
                    if self.nombre_hoja in libro_verificacion.sheetnames
                    else libro_verificacion.active
                )
                if hoja_verificacion is None:
                    raise ValueError(
                        "No se pudo verificar la plantilla después de la edición."
                    )
                return [
                    fila
                    for fila in range(
                        self._fila_inicio_datos,
                        hoja_verificacion.max_row + 1,
                    )
                    if hoja_verificacion.cell(
                        fila, self._col_codigo
                    ).value is not None
                    and str(
                        hoja_verificacion.cell(fila, self._col_nombre).value or ""
                    ).strip()
                ]
            finally:
                libro_verificacion.close()

        def restaurar_plantilla() -> None:
            if not plantilla_actualizada:
                return
            with tempfile.NamedTemporaryFile(
                suffix=self.ruta_excel_base.suffix,
                dir=self.ruta_excel_base.parent,
                delete=False,
            ) as temporal:
                temporal.write(contenido_original)
                ruta_temporal = Path(temporal.name)
            try:
                os.replace(ruta_temporal, self.ruta_excel_base)
            finally:
                if ruta_temporal.exists():
                    ruta_temporal.unlink()

        self.product_store.update_product(
            old_code=codigo_actual,
            new_code=codigo_nuevo,
            name=nombre,
            stock=stock,
            synchronize_template=sincronizar_plantilla,
            restore_template=restaurar_plantilla,
            stock_minimo=stock_minimo,
            stock_maximo=stock_maximo,
        )
        self._cargar_catalogo()

    def cambiar_estado_producto(self, codigo: str, activo: bool) -> None:
        self.product_store.set_product_active(normalizar_codigo(codigo), activo)
        self._cargar_catalogo()

    def agregar_producto_en_posicion(
        self,
        codigo_nuevo: str,
        nombre_nuevo: str,
        codigo_arriba: str,
        stock_inicial: Optional[int] = None
    ) -> Tuple[bool, str, Optional[int]]:
        """
        Inserta un nuevo producto en SQLite inmediatamente debajo del código
        indicado por codigo_arriba, conservando el orden del catálogo.

        Retorna: (exito: bool, mensaje: str, fila_insertada: Optional[int])
        """
        cod_nuevo_norm = normalizar_codigo(codigo_nuevo)
        cod_arriba_norm = normalizar_codigo(codigo_arriba)

        # 1. Validar que el nuevo código no exista ya
        if cod_nuevo_norm in self.catalogo:
            fila_exist, nombre_exist, _ = self.catalogo[cod_nuevo_norm]
            return False, f"El codigo '{cod_nuevo_norm}' ya existe en la fila {fila_exist} ('{nombre_exist}').", None

        # 2. Localizar la fila del producto superior de referencia
        if cod_arriba_norm not in self.catalogo:
            return False, f"No se encontro el producto de referencia con codigo '{cod_arriba_norm}' en el catalogo.", None

        contenido_original = self.ruta_excel_base.read_bytes()
        respaldo = self.ruta_excel_base.with_name(
            f"{self.ruta_excel_base.stem}_backup{self.ruta_excel_base.suffix}"
        )
        plantilla_actualizada = False

        def sincronizar_plantilla(orden: int) -> list[int]:
            nonlocal plantilla_actualizada
            shutil.copy2(self.ruta_excel_base, respaldo)
            libro = openpyxl.load_workbook(str(self.ruta_excel_base))
            try:
                hoja = (
                    libro[self.nombre_hoja]
                    if self.nombre_hoja in libro.sheetnames
                    else libro.active
                )
                if hoja is None:
                    raise ValueError("No se encontró una hoja para actualizar.")
                fila_referencia = next(
                    (
                        fila
                        for fila in range(
                            self._fila_inicio_datos, hoja.max_row + 1
                        )
                        if normalizar_codigo(
                            str(
                                hoja.cell(fila, self._col_codigo).value or ""
                            )
                        )
                        == cod_arriba_norm
                    ),
                    None,
                )
                if fila_referencia is None:
                    raise ValueError(
                        f"No se encontró el código de referencia "
                        f"'{cod_arriba_norm}' en Productos.xlsx."
                    )

                fila_nueva = fila_referencia + 1
                hoja.insert_rows(fila_nueva, 1)
                for columna in range(1, hoja.max_column + 1):
                    origen = hoja.cell(fila_referencia, columna)
                    destino = hoja.cell(fila_nueva, columna)
                    if origen.has_style:
                        destino._style = copy(origen._style)
                    if origen.number_format:
                        destino.number_format = origen.number_format
                hoja.row_dimensions[fila_nueva].height = (
                    hoja.row_dimensions[fila_referencia].height
                )

                codigo_excel = (
                    int(cod_nuevo_norm)
                    if cod_nuevo_norm.isdigit()
                    else cod_nuevo_norm
                )
                hoja.cell(fila_nueva, self._col_codigo, codigo_excel)
                hoja.cell(
                    fila_nueva,
                    self._col_nombre,
                    str(nombre_nuevo).strip(),
                )
                stock = stock_inicial or 0
                if self._col_stock_final is not None:
                    columna_inicial = self._col_stock_inicial or self._col_stock
                    hoja.cell(fila_nueva, columna_inicial, stock)
                    if self._col_descuento is not None:
                        hoja.cell(fila_nueva, self._col_descuento, None)
                    hoja.cell(fila_nueva, self._col_stock_final, stock)
                else:
                    hoja.cell(fila_nueva, self._col_stock, stock)

                with tempfile.NamedTemporaryFile(
                    suffix=self.ruta_excel_base.suffix,
                    dir=self.ruta_excel_base.parent,
                    delete=False,
                ) as temporal:
                    ruta_temporal = Path(temporal.name)
                try:
                    libro.save(ruta_temporal)
                finally:
                    libro.close()
                try:
                    os.replace(ruta_temporal, self.ruta_excel_base)
                    plantilla_actualizada = True
                finally:
                    if ruta_temporal.exists():
                        ruta_temporal.unlink()
            except Exception:
                libro.close()
                raise

            libro_verificacion = openpyxl.load_workbook(
                str(self.ruta_excel_base), data_only=True, read_only=True
            )
            try:
                hoja_verificacion = (
                    libro_verificacion[self.nombre_hoja]
                    if self.nombre_hoja in libro_verificacion.sheetnames
                    else libro_verificacion.active
                )
                if hoja_verificacion is None:
                    raise ValueError(
                        "No se pudo verificar la plantilla después del alta."
                    )
                filas_plantilla = [
                    fila
                    for fila in range(
                        self._fila_inicio_datos,
                        hoja_verificacion.max_row + 1,
                    )
                    if hoja_verificacion.cell(
                        fila, self._col_codigo
                    ).value is not None
                    and str(
                        hoja_verificacion.cell(
                            fila, self._col_nombre
                        ).value
                        or ""
                    ).strip()
                ]
                if (
                    orden >= len(filas_plantilla)
                    or filas_plantilla[orden] != fila_nueva
                ):
                    raise RuntimeError(
                        "El orden insertado en Productos.xlsx no coincide "
                        "con el orden de SQLite."
                    )
                return filas_plantilla
            finally:
                libro_verificacion.close()

        def restaurar_plantilla() -> None:
            if not plantilla_actualizada:
                return
            with tempfile.NamedTemporaryFile(
                suffix=self.ruta_excel_base.suffix,
                dir=self.ruta_excel_base.parent,
                delete=False,
            ) as temporal:
                temporal.write(contenido_original)
                ruta_temporal = Path(temporal.name)
            try:
                os.replace(ruta_temporal, self.ruta_excel_base)
            finally:
                if ruta_temporal.exists():
                    ruta_temporal.unlink()

        try:
            orden, nombre_arriba = self.product_store.add_product_after(
                code=cod_nuevo_norm,
                name=str(nombre_nuevo).strip(),
                reference_code=cod_arriba_norm,
                stock=stock_inicial or 0,
                synchronize_template=sincronizar_plantilla,
                restore_template=restaurar_plantilla,
            )
        except (ValueError, OSError, RuntimeError) as error:
            return False, str(error), None

        self._cargar_catalogo()

        fila_nueva = orden + self._fila_inicio_datos
        return True, (
            f"Producto '{nombre_nuevo}' [{cod_nuevo_norm}] insertado en la "
            f"posición {fila_nueva} del catálogo (debajo de [{cod_arriba_norm}] "
            f"{nombre_arriba})."
        ), fila_nueva

    def aplicar_descuentos(
        self,
        items_consolidados: Dict[str, ItemConsolidado],
        carpeta_destino: str | Path,
        pedidos_procesados: Optional[List[ResultadoRemito]] = None,
        guardar_archivo: bool = True,
        prefijo_archivo: str = "Descuento",
    ) -> ReporteStock:
        """Genera un informe con el diseño del catálogo, sin modificarlo."""
        carpeta_destino = Path(carpeta_destino)
        fecha_generacion = datetime.now()
        fecha_hora_str = fecha_generacion.strftime("%d-%m-%Y_%H-%M-%S")
        nombre_archivo = f"{prefijo_archivo}_{fecha_hora_str}.xls" if prefijo_archivo != "Descuento" else f"Descuento{fecha_hora_str}.xls"
        ruta_guardado = carpeta_destino / nombre_archivo
        ruta_pdf = ruta_guardado.with_suffix(".pdf")

        descuentos: List[RegistroDescuento] = []
        for _, cod, nombre, stock_ant in self.productos_base:
            item = items_consolidados.get(cod)
            if item is None:
                continue
            descuentos.append(RegistroDescuento(
                codigo=cod,
                nombre_producto=nombre,
                fila_excel=4 + len(descuentos),
                stock_anterior=stock_ant,
                cantidad_descontada=item.cantidad_total,
                stock_nuevo=int(stock_ant - item.cantidad_total),
            ))

        filas_productos = []
        for fila_original, cod, nombre, _ in self.filas_catalogo:
            item = items_consolidados.get(cod)
            filas_productos.append((
                fila_original,
                cod,
                item.cantidad_total if item else None,
                nombre,
                -item.cantidad_total if item else None,
            ))

        codigos_descontados = {registro.codigo for registro in descuentos}
        pedidos_descontados = [
            documento
            for documento in (pedidos_procesados or [])
            if not documento.items
            or any(
                normalizar_codigo(item.codigo) in codigos_descontados
                for item in documento.items
            )
        ]
        productos_sin_coincidencia = [
            (
                cod,
                ", ".join(item.remitos_origen),
                item.descripcion,
                "Código no encontrado en el catálogo SQLite",
            )
            for cod, item in items_consolidados.items()
            if cod not in self.catalogo
        ]
        productos_sin_coincidencia.extend(
            (
                "",
                documento.archivo,
                descripcion,
                motivo,
            )
            for documento in (pedidos_procesados or [])
            for descripcion, motivo in documento.descripciones_sin_coincidencia
        )

        wb_salida, ws_salida = self._crear_informe_desde_plantilla(
            len(descuentos),
            sum(descuento.cantidad_descontada for descuento in descuentos),
            fecha_generacion.strftime("%d/%m/%y"),
            filas_productos,
            pedidos_descontados,
            productos_sin_coincidencia,
        )

        if guardar_archivo:
            carpeta_destino.mkdir(parents=True, exist_ok=True)
            self._guardar_informe_xls(wb_salida, ruta_guardado)
            self._guardar_informe_pdf(wb_salida, ruta_pdf)
            archivo_final = str(ruta_guardado)
            archivo_pdf = str(ruta_pdf)
        else:
            archivo_final = ""
            archivo_pdf = None

        return ReporteStock(
            descuentos_aplicados=descuentos,
            archivo_guardado=archivo_final,
            archivo_pdf=archivo_pdf,
        )

    def _crear_informe_desde_plantilla(
        self,
        cantidad_descuentos: int,
        unidades_descontadas: int,
        fecha_generacion: str,
        filas_productos: List[Tuple[int, str, Optional[int], str, Optional[int]]],
        pedidos_descontados: List[ResultadoRemito],
        productos_sin_coincidencia: List[Tuple[str, str, str, str]],
    ):
        libro_plantilla = openpyxl.load_workbook(self.ruta_excel_base)
        hoja_plantilla = (
            libro_plantilla[self.nombre_hoja]
            if self.nombre_hoja in libro_plantilla.sheetnames
            else libro_plantilla.active
        )
        libro_salida = openpyxl.Workbook()
        hoja_salida = libro_salida.active
        hoja_salida.title = hoja_plantilla.title[:31]
        hoja_salida.sheet_view.showGridLines = hoja_plantilla.sheet_view.showGridLines
        hoja_salida.page_setup.orientation = "portrait"
        hoja_salida.page_setup.paperSize = hoja_salida.PAPERSIZE_A4
        hoja_salida.page_setup.fitToWidth = 1
        hoja_salida.page_setup.fitToHeight = 0
        hoja_salida.sheet_properties.pageSetUpPr.fitToPage = True
        margen_cm = 2 / 2.54
        hoja_salida.page_margins.left = margen_cm
        hoja_salida.page_margins.right = margen_cm
        hoja_salida.page_margins.top = margen_cm
        hoja_salida.page_margins.bottom = margen_cm
        hoja_salida.freeze_panes = "A4"

        def copiar_estilo(origen, destino) -> None:
            if not origen.has_style:
                return
            destino.font = copy(origen.font)
            destino.fill = copy(origen.fill)
            destino.border = copy(origen.border)
            destino.alignment = copy(origen.alignment)
            destino.number_format = origen.number_format
            destino.protection = copy(origen.protection)

        def aplicar_encabezado(
            fila: int, color_fondo: str, cantidad_columnas: int = 4
        ) -> None:
            for columna in range(1, cantidad_columnas + 1):
                celda = hoja_salida.cell(fila, columna)
                celda.fill = PatternFill("solid", fgColor=color_fondo)
                celda.font = Font(
                    name="Arial",
                    size=9,
                    bold=True,
                    color="17365D",
                )
                celda.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                celda.border = Border(
                    bottom=Side(style="thin", color="D9E2F3")
                )

        def aplicar_titulo_seccion(fila: int, color_fondo: str) -> None:
            celda = hoja_salida.cell(fila, 1)
            celda.fill = PatternFill("solid", fgColor=color_fondo)
            celda.font = Font(name="Arial", size=11, bold=True, color="17365D")
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        columnas_plantilla = (
            self._col_codigo,
            self._col_nombre,
            self._col_descuento or self._col_stock,
        )
        for columna_salida, columna_plantilla in enumerate(
            columnas_plantilla, start=1
        ):
            letra_salida = openpyxl.utils.get_column_letter(columna_salida)
            letra_plantilla = openpyxl.utils.get_column_letter(columna_plantilla)
            dimension = hoja_plantilla.column_dimensions[letra_plantilla]
            ancho = dimension.width
            if columna_salida == 3:
                ancho = max(ancho or 0, 22)
            hoja_salida.column_dimensions[letra_salida].width = ancho
            hoja_salida.column_dimensions[letra_salida].hidden = dimension.hidden

        for fila in range(1, 4):
            hoja_salida.row_dimensions[fila].height = (
                hoja_plantilla.row_dimensions[fila].height
            )
            for columna_salida, columna_plantilla in enumerate(
                columnas_plantilla, start=1
            ):
                origen = hoja_plantilla.cell(fila, columna_plantilla)
                destino = hoja_salida.cell(fila, columna_salida)
                copiar_estilo(origen, destino)

        for encabezado in ("A1:A2", "B1:B2", "C1:D2"):
            hoja_salida.merge_cells(encabezado)
        hoja_salida["A1"] = (
            hoja_plantilla.cell(1, self._col_codigo).value or "CODIGO"
        )
        hoja_salida["B1"] = (
            hoja_plantilla.cell(1, self._col_nombre).value or "PRODUCTOS"
        )
        hoja_salida["C1"] = f"FECHA: {fecha_generacion}"
        aplicar_encabezado(1, "D9EAF7", 3)
        hoja_salida["D1"].fill = PatternFill("solid", fgColor="D9EAF7")
        hoja_salida["D1"].border = Border(
            bottom=Side(style="thin", color="D9E2F3")
        )

        encabezado_productos = hoja_plantilla.cell(3, self._col_nombre).value
        hoja_salida["B3"] = (
            encabezado_productos
            if self._col_stock_inicial is not None and encabezado_productos
            else "PRODUCTOS"
        )
        hoja_salida["C3"] = "DESCUENTO"
        hoja_salida["C3"].number_format = "@"
        aplicar_encabezado(3, "D9EAF7", 3)
        hoja_salida.row_dimensions[1].height = 48
        hoja_salida.row_dimensions[2].height = 10
        hoja_salida.row_dimensions[3].height = 24

        fila_estilo_producto = min(
            max(self._fila_inicio_datos, 1),
            max(hoja_plantilla.max_row, 1),
        )
        for fila_salida, producto in enumerate(filas_productos, start=4):
            fila_original, codigo, unidades, nombre, descuento = producto
            fila_estilo = (
                fila_original
                if 0 < fila_original <= hoja_plantilla.max_row
                else fila_estilo_producto
            )
            hoja_salida.row_dimensions[fila_salida].height = (
                hoja_plantilla.row_dimensions[fila_estilo].height
            )
            for columna_salida, columna_plantilla in enumerate(
                columnas_plantilla, start=1
            ):
                origen = hoja_plantilla.cell(fila_estilo, columna_plantilla)
                destino = hoja_salida.cell(fila_salida, columna_salida)
                copiar_estilo(origen, destino)
            codigo_salida = int(codigo) if codigo.isdigit() else codigo
            hoja_salida.cell(fila_salida, 1, codigo_salida)
            hoja_salida.cell(fila_salida, 2, nombre)
            if descuento is not None:
                hoja_salida.cell(fila_salida, 3, descuento).number_format = "0"

        fila_actual = 4 + len(filas_productos)
        hoja_salida.cell(fila_actual, 2, "TOTAL CÓDIGOS DE PRODUCTO DESCONTADOS:")
        hoja_salida.cell(fila_actual, 3, cantidad_descuentos)
        self._copiar_estilo_fila(
            hoja_plantilla, hoja_salida, 3, fila_actual, 3
        )
        for columna in range(1, 4):
            celda = hoja_salida.cell(fila_actual, columna)
            celda.fill = PatternFill("solid", fgColor="D9EAF7")
            celda.font = Font(name="Arial", size=10, bold=True, color="17365D")
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        fila_actual += 1
        hoja_salida.cell(fila_actual, 2, "TOTAL UNIDADES DESCONTADAS:")
        hoja_salida.cell(fila_actual, 3, unidades_descontadas)
        self._copiar_estilo_fila(
            hoja_plantilla, hoja_salida, 3, fila_actual, 3
        )
        for columna in range(1, 4):
            celda = hoja_salida.cell(fila_actual, columna)
            celda.fill = PatternFill("solid", fgColor="D9EAF7")
            celda.font = Font(name="Arial", size=10, bold=True, color="17365D")
            celda.alignment = Alignment(vertical="center", wrap_text=True)
        fila_actual += 1
        hoja_salida.merge_cells(
            start_row=fila_actual,
            start_column=1,
            end_row=fila_actual,
            end_column=4,
        )
        hoja_salida.cell(fila_actual, 1, "PEDIDOS CON PRODUCTOS DESCONTADOS")
        self._copiar_estilo_fila(
            hoja_plantilla, hoja_salida, 3, fila_actual, 4
        )
        aplicar_titulo_seccion(fila_actual, "EAF0F6")
        fila_actual += 1
        for columna, titulo in enumerate(
            ("TIPO", "CLIENTE", "N° REMITO / PEDIDO"), start=1
        ):
            hoja_salida.cell(fila_actual, columna, titulo)
        self._copiar_estilo_fila(
            hoja_plantilla, hoja_salida, 3, fila_actual, 4
        )
        aplicar_encabezado(fila_actual, "D9EAF7")
        if pedidos_descontados:
            for documento in pedidos_descontados:
                fila_actual += 1
                if documento.archivo.lower().endswith(".pdf"):
                    tipo = "Remito Mayorista"
                elif documento.archivo.lower().endswith(".docx"):
                    tipo = "Mercado Libre"
                else:
                    tipo = "Canje"
                for columna, valor in enumerate(
                    (tipo, documento.cliente, str(documento.numero_remito)),
                    start=1,
                ):
                    hoja_salida.cell(fila_actual, columna, valor)
                self._copiar_estilo_fila(
                    hoja_plantilla, hoja_salida, fila_estilo_producto,
                    fila_actual, 4,
                )
        else:
            fila_actual += 1
            hoja_salida.merge_cells(
                start_row=fila_actual,
                start_column=1,
                end_row=fila_actual,
                end_column=4,
            )
            hoja_salida.cell(fila_actual, 1, "No hubo pedidos con productos descontados.")
            self._copiar_estilo_fila(
                hoja_plantilla, hoja_salida, fila_estilo_producto,
                fila_actual, 4,
            )
        fila_actual += 1

        if productos_sin_coincidencia:
            hoja_salida.merge_cells(
                start_row=fila_actual,
                start_column=1,
                end_row=fila_actual,
                end_column=4,
            )
            hoja_salida.cell(
                fila_actual,
                1,
                "PRODUCTOS NO DESCONTADOS - SIN COINCIDENCIA",
            )
            self._copiar_estilo_fila(
                hoja_plantilla, hoja_salida, 3, fila_actual, 4
            )
            aplicar_titulo_seccion(fila_actual, "FCE4D6")
            fila_actual += 1
            for columna, titulo in enumerate(
                ("CÓDIGO", "ARCHIVO / PEDIDO", "PRODUCTO / DESCRIPCIÓN", "MOTIVO"),
                start=1,
            ):
                hoja_salida.cell(fila_actual, columna, titulo)
            self._copiar_estilo_fila(
                hoja_plantilla, hoja_salida, 3, fila_actual, 4
            )
            aplicar_encabezado(fila_actual, "D9EAF7")
            for codigo, archivo, descripcion, motivo in productos_sin_coincidencia:
                fila_actual += 1
                hoja_salida.cell(fila_actual, 1, codigo)
                hoja_salida.cell(fila_actual, 2, archivo)
                hoja_salida.cell(fila_actual, 3, descripcion)
                hoja_salida.cell(fila_actual, 4, motivo)
                self._copiar_estilo_fila(
                    hoja_plantilla, hoja_salida, fila_estilo_producto,
                    fila_actual, 4,
                )

        libro_plantilla.close()
        hoja_salida.print_area = f"A1:D{max(3, fila_actual)}"
        return libro_salida, hoja_salida

    def _copiar_estilo_fila(
        self,
        hoja_plantilla,
        hoja_salida,
        fila_plantilla: int,
        fila_salida: int,
        cantidad_columnas: int,
    ) -> None:
        columnas_plantilla = (
            self._col_codigo,
            self._col_nombre,
            self._col_descuento or self._col_stock,
            self._col_stock_final or self._col_stock,
        )
        copiar_estilo_fila(
            hoja_plantilla=hoja_plantilla,
            hoja_salida=hoja_salida,
            fila_plantilla=fila_plantilla,
            fila_salida=fila_salida,
            cantidad_columnas=cantidad_columnas,
            columnas_plantilla=columnas_plantilla,
        )

    @staticmethod
    def _guardar_informe_xls(libro: openpyxl.Workbook, ruta: Path) -> None:
        guardar_informe_xls(libro, ruta)

    @staticmethod
    def _guardar_informe_pdf(libro: openpyxl.Workbook, ruta: Path) -> None:
        guardar_informe_pdf(libro, ruta)

    def sincronizar_stocks(self, stocks: Dict[str, int]) -> int:
        """Actualiza el stock en la base SQLite y sincroniza Productos.xlsx."""
        if not stocks:
            return 0

        stocks_normalizados = {
            normalizar_codigo(k): int(v) for k, v in stocks.items()
        }

        contenido_original = self.ruta_excel_base.read_bytes()
        respaldo = self.ruta_excel_base.with_name(
            f"{self.ruta_excel_base.stem}_backup{self.ruta_excel_base.suffix}"
        )
        plantilla_actualizada = False

        def sincronizar_plantilla() -> list[int]:
            nonlocal plantilla_actualizada
            shutil.copy2(self.ruta_excel_base, respaldo)
            libro = openpyxl.load_workbook(str(self.ruta_excel_base))
            try:
                hoja = (
                    libro[self.nombre_hoja]
                    if self.nombre_hoja in libro.sheetnames
                    else libro.active
                )
                if hoja is None:
                    raise ValueError("No se encontró una hoja para actualizar.")

                for fila in range(self._fila_inicio_datos, hoja.max_row + 1):
                    val_cod = hoja.cell(fila, self._col_codigo).value
                    if val_cod is None:
                        continue
                    cod = normalizar_codigo(str(val_cod))
                    if cod in stocks_normalizados:
                        nuevo_stock = stocks_normalizados[cod]
                        if self._col_stock_final is not None:
                            columna_inicial = (
                                self._col_stock_inicial or self._col_stock
                            )
                            hoja.cell(fila, columna_inicial, nuevo_stock)
                            if self._col_descuento is not None:
                                hoja.cell(fila, self._col_descuento, None)
                            hoja.cell(fila, self._col_stock_final, nuevo_stock)
                        else:
                            hoja.cell(fila, self._col_stock, nuevo_stock)

                with tempfile.NamedTemporaryFile(
                    suffix=self.ruta_excel_base.suffix,
                    dir=self.ruta_excel_base.parent,
                    delete=False,
                ) as temporal:
                    ruta_temporal = Path(temporal.name)
                try:
                    libro.save(ruta_temporal)
                finally:
                    libro.close()
                try:
                    os.replace(ruta_temporal, self.ruta_excel_base)
                    plantilla_actualizada = True
                finally:
                    if ruta_temporal.exists():
                        ruta_temporal.unlink()

            except Exception:
                libro.close()
                raise

            libro_verificacion = openpyxl.load_workbook(
                str(self.ruta_excel_base), data_only=True, read_only=True
            )
            try:
                hoja_verificacion = (
                    libro_verificacion[self.nombre_hoja]
                    if self.nombre_hoja in libro_verificacion.sheetnames
                    else libro_verificacion.active
                )
                if hoja_verificacion is None:
                    raise ValueError(
                        "No se pudo verificar la plantilla después de la actualización de stock."
                    )
                return [
                    fila
                    for fila in range(
                        self._fila_inicio_datos,
                        hoja_verificacion.max_row + 1,
                    )
                    if hoja_verificacion.cell(
                        fila, self._col_codigo
                    ).value is not None
                    and str(
                        hoja_verificacion.cell(fila, self._col_nombre).value or ""
                    ).strip()
                ]
            finally:
                libro_verificacion.close()

        def restaurar_plantilla() -> None:
            if not plantilla_actualizada:
                return
            with tempfile.NamedTemporaryFile(
                suffix=self.ruta_excel_base.suffix,
                dir=self.ruta_excel_base.parent,
                delete=False,
            ) as temporal:
                temporal.write(contenido_original)
                ruta_temporal = Path(temporal.name)
            try:
                os.replace(ruta_temporal, self.ruta_excel_base)
            finally:
                if ruta_temporal.exists():
                    ruta_temporal.unlink()

        filas_actualizadas = self.product_store.update_stocks(
            stocks=stocks_normalizados,
            synchronize_template=sincronizar_plantilla,
            restore_template=restaurar_plantilla,
        )
        self._cargar_catalogo()
        return filas_actualizadas

    def registrar_movimientos_historial(
        self,
        stocks_delta: Dict[str, int],
        tipo: str,
    ) -> None:
        """Persiste movimientos en stock_movements para estadísticas.

        stocks_delta: {codigo: cantidad} — positivo=ingreso, negativo=descuento.
        tipo: 'descuento' | 'ingreso_fabrica' | 'ajuste_manual'
        """
        movimientos = []
        for cod, delta in stocks_delta.items():
            nombre = self.catalogo.get(cod, (0, cod, 0))[1]
            movimientos.append((cod, nombre, delta, tipo))
        self.product_store.register_movements(movimientos)

    def consultar_movimientos(
        self,
        fecha_desde: str,
        fecha_hasta: str,
        tipo: Optional[str] = None,
        top_n: Optional[int] = None,
    ) -> List:
        """Consulta historial agrupado. Retorna lista de sqlite3.Row."""
        hasta_inclusive = fecha_hasta if "T" in fecha_hasta else fecha_hasta + "T23:59:59"
        desde_inclusive = fecha_desde if "T" in fecha_desde else fecha_desde + "T00:00:00"
        return self.product_store.query_movements(
            date_from=desde_inclusive,
            date_to=hasta_inclusive,
            movement_type=tipo,
            top_n=top_n,
        )

    def registrar_ingreso_fabrica(
        self,
        ingresos: Dict[str, int],
        sobrescribir: bool = False,
    ) -> Dict[str, Tuple[int, int]]:
        """Aplica un ingreso, descuento o restablecimiento de stock.

        - sobrescribir=False (por defecto): suma la cantidad al stock actual.
          Si la cantidad es negativa descuenta ese monto (el resultado puede quedar negativo).
        - sobrescribir=True: establece el stock exactamente al valor indicado (mínimo 0).

        Retorna un diccionario {codigo: (stock_anterior, stock_nuevo)}.
        """
        if not ingresos:
            return {}

        resumen: Dict[str, Tuple[int, int]] = {}
        actualizaciones: Dict[str, int] = {}

        for codigo_raw, cantidad in ingresos.items():
            cod = normalizar_codigo(codigo_raw)
            if cod not in self.catalogo:
                raise ValueError(
                    f"El código '{cod}' no se encuentra en el catálogo de productos."
                )
            _, _, stock_actual = self.catalogo[cod]
            if sobrescribir:
                # Ajuste directo: fija el valor indicado (mínimo 0 en modo sobrescribir)
                nuevo_stock = max(0, int(cantidad))
            else:
                # Suma/resta libre: permite negativos para reflejar sobredemanda o descuentos manuales
                nuevo_stock = int(stock_actual) + int(cantidad)
            actualizaciones[cod] = nuevo_stock
            resumen[cod] = (stock_actual, nuevo_stock)

        self.sincronizar_stocks(actualizaciones)
        return resumen

    def actualizar_stock_base(
        self,
        descuentos: List[RegistroDescuento],
        sincronizar_excel: bool = False,
        nombre_archivo: str = "Descuento",
        usuario: str = "sistema",
    ) -> int:
        """Persiste en SQLite los saldos nuevos de cada producto descontado.
        
        Si sincronizar_excel es True, también sincroniza el stock en Productos.xlsx.
        """
        existentes = [d for d in descuentos if not d.es_nuevo_agregado]
        if not existentes:
            return 0

        actualizaciones = {
            descuento.codigo: descuento.stock_nuevo for descuento in existentes
        }

        # Registrar lote en despachos y despacho_detalles
        total_unidades = sum(d.cantidad_descontada for d in existentes)
        try:
            detalles = [
                (d.codigo, d.cantidad_descontada, d.stock_anterior, d.stock_nuevo)
                for d in existentes
            ]
            self.product_store.record_dispatch(
                tipo_origen="Descuento",
                nombre_archivo=nombre_archivo,
                total_articulos=total_unidades,
                detalles=detalles,
                usuario=usuario,
            )
        except Exception as e:
            print(f"[AVISO] No se pudo registrar detalle de despacho: {e}")

        # Si el descuento lo realiza un operario (o usuario no admin/sistema), notificar al administrador
        if usuario and usuario.lower() not in ("sistema", "administrador", "admin"):
            try:
                self.product_store.create_admin_notification(
                    title=f"Descuento procesado por operario '{usuario}'",
                    message=(
                        f"El operario '{usuario}' ejecutó un descuento de stock de {total_unidades} unidades "
                        f"en {len(existentes)} producto(s) desde {nombre_archivo}."
                    ),
                    operator_username=usuario,
                    notification_type="descuento",
                    details=f"Archivo: {nombre_archivo} | Unidades: {total_unidades} | Productos: {len(existentes)}",
                )
            except Exception as notif_err:
                print(f"[AVISO] No se pudo generar notificación para administrador: {notif_err}")

        if sincronizar_excel:
            return self.sincronizar_stocks(actualizaciones)

        filas_actualizadas = self.product_store.update_stocks(actualizaciones)
        self._cargar_catalogo()
        return filas_actualizadas

    def obtener_notificaciones_admin(self, solo_no_leidas: bool = False) -> list:
        return self.product_store.list_admin_notifications(unread_only=solo_no_leidas)

    def contar_notificaciones_no_leidas(self) -> int:
        return self.product_store.get_unread_notifications_count()

    def marcar_notificaciones_leidas(self, notificacion_id: int | None = None) -> None:
        if notificacion_id is not None:
            self.product_store.mark_notification_as_read(notificacion_id)
        else:
            self.product_store.mark_all_notifications_as_read()

    def descontar_mercado_full(
        self,
        descuentos: Dict[str, int],
        nombre_archivo: str = "MercadoFull",
        sincronizar_excel: bool = True,
    ) -> Dict[str, Tuple[int, int]]:
        """Descuenta las unidades enviadas a Mercado Envíos Full del stock actual.

        descuentos: Dict[codigo, cantidad_a_descontar]
        Retorna: Dict[codigo, (stock_anterior, stock_nuevo)]
        """
        if not descuentos:
            return {}

        resumen: Dict[str, Tuple[int, int]] = {}
        actualizaciones: Dict[str, int] = {}
        detalles_despacho = []

        for codigo_raw, cant_descontar in descuentos.items():
            cod = normalizar_codigo(codigo_raw)
            if cod not in self.catalogo:
                raise ValueError(
                    f"El código '{cod}' no se encuentra en el catálogo de productos."
                )
            cant = int(cant_descontar)
            if cant <= 0:
                continue

            _, _, stock_actual = self.catalogo[cod]
            nuevo_stock = int(stock_actual) - cant
            actualizaciones[cod] = nuevo_stock
            resumen[cod] = (stock_actual, nuevo_stock)
            detalles_despacho.append((cod, cant, stock_actual, nuevo_stock))

        if not actualizaciones:
            return {}

        try:
            self.product_store.record_dispatch(
                tipo_origen="MercadoFull",
                nombre_archivo=nombre_archivo,
                total_articulos=sum(cant for _, cant, _, _ in detalles_despacho),
                detalles=detalles_despacho,
            )
        except Exception as e:
            print(f"[AVISO] No se pudo registrar despacho de Mercado Full: {e}")

        if sincronizar_excel:
            self.sincronizar_stocks(actualizaciones)
        else:
            self.product_store.update_stocks(actualizaciones)
            self._cargar_catalogo()

        return resumen

    def obtener_analisis_reposicion(self, solo_criticos: bool = True) -> List:
        """Consulta la vista de la base de datos v_analisis_reposicion."""
        return self.product_store.get_replenishment_analysis(only_critical=solo_criticos)

    def generar_pedido_sugerido_fabrica(
        self,
        ruta_salida: Path | str | None = None,
        solo_criticos: bool = True,
    ) -> Path:
        """Genera un archivo Excel con el pedido sugerido a fábrica.

        Calcula la cantidad sugerida como (stock_maximo - stock) para los productos
        que requieran reposición según sus umbrales configurados.
        """
        analisis = self.obtener_analisis_reposicion(solo_criticos=solo_criticos)
        return generar_pedido_sugerido_fabrica_excel(
            analisis=analisis,
            ruta_excel_base=self.ruta_excel_base,
            ruta_salida=ruta_salida,
        )

    def comparar_stock_fisico(self, conteo_fisico: Dict[str, int]) -> List[Dict]:
        """Compara un diccionario {codigo: stock_fisico} con el stock digital actual.

        Retorna una lista de diccionarios con:
            code, name, stock_digital, stock_fisico, diferencia, estado, en_catalogo
        """
        return comparar_stock_fisico_dict(self.catalogo, conteo_fisico)

    def procesar_stock_fisico(
        self,
        conteo_fisico: Dict[str, int],
        usuario: str = "operador",
        nombre_archivo: str = "StockFisico.xlsx",
        ruta_salida: Path | str | None = None,
    ) -> Tuple[Path, List[Dict]]:
        """Actualiza el stock físico en SQLite, compara con el stock digital,
        genera el documento Excel unificado de comparativa y diferencias para el administrador,
        y crea la notificación correspondiente indicando el operario responsable.
        """
        comparativa = self.comparar_stock_fisico(conteo_fisico)
        ruta_doc = self.generar_documento_comparativo_stock_fisico(
            comparativa=comparativa,
            usuario=usuario,
            nombre_archivo_origen=nombre_archivo,
            ruta_salida=ruta_salida,
        )

        detalles_db = [
            (
                str(item["code"]),
                str(item["name"]),
                int(item["stock_digital"]),
                int(item["stock_fisico"]),
                int(item["diferencia"]),
                str(item["estado"]),
            )
            for item in comparativa
        ]
        self.product_store.record_physical_inventory(
            usuario=usuario,
            nombre_archivo=nombre_archivo,
            reporte_unificado=str(ruta_doc),
            detalles=detalles_db,
        )

        total_relevados = len(comparativa)
        total_coinciden = sum(1 for i in comparativa if i["diferencia"] == 0 and i["en_catalogo"])
        total_faltantes = sum(1 for i in comparativa if i["diferencia"] < 0 and i["en_catalogo"])
        total_sobrantes = sum(1 for i in comparativa if i["diferencia"] > 0 and i["en_catalogo"])
        total_diferencias = total_faltantes + total_sobrantes

        try:
            self.product_store.create_admin_notification(
                title=f"Inventario físico cargado por operario '{usuario}'",
                message=(
                    f"El operario '{usuario}' cargó un inventario físico desde '{nombre_archivo}' "
                    f"({total_relevados} productos relevados: {total_coinciden} coinciden, "
                    f"{total_faltantes} con faltante, {total_sobrantes} con sobrante). "
                    f"Documento unificado: {ruta_doc.name}."
                ),
                operator_username=usuario,
                notification_type="stock_fisico",
                details=f"Archivo: {ruta_doc.name} | Relevados: {total_relevados} | Diferencias: {total_diferencias}",
            )
        except Exception as notif_err:
            print(f"[AVISO] No se pudo crear notificación de stock físico: {notif_err}")

        return ruta_doc, comparativa

    def generar_documento_comparativo_stock_fisico(
        self,
        comparativa: List[Dict],
        usuario: str = "operador",
        nombre_archivo_origen: str = "",
        ruta_salida: Path | str | None = None,
    ) -> Path:
        """Genera un archivo Excel unificado con la comparativa entre Stock Digital y Stock Físico
        y el detalle de diferencias para el administrador.
        """
        return generar_documento_comparativo_excel(
            comparativa=comparativa,
            ruta_excel_base=self.ruta_excel_base,
            product_store=self.product_store,
            usuario=usuario,
            nombre_archivo_origen=nombre_archivo_origen,
            ruta_salida=ruta_salida,
        )

    def obtener_ultima_comparativa_fisica(self):
        return self.product_store.get_latest_physical_inventory()

    def listar_inventarios_fisicos(self, limit: int = 20):
        return self.product_store.list_physical_inventories(limit=limit)

    def obtener_detalle_inventario_fisico(self, inventario_id: int):
        return self.product_store.get_physical_inventory_details(inventario_id)
