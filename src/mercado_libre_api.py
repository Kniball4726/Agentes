"""Módulo Conector Oficial con la API REST de Mercado Libre para StockAgent.

Permite:
1. Configurar credenciales OAuth (App ID, Client Secret, Access Token, Refresh Token).
2. Probar la conexión en tiempo real con la API de Mercado Libre (`/users/me`).
3. Renovar automáticamente el Access Token vía OAuth cuando expire.
4. Consultar y sincronizar desatendidamente órdenes recientes pagadas (`/orders/search`).
5. Emparejar artículos por SKU del vendedor (100% de precisión) o por título/descripción (fuzzy matching).
6. Descontar stock de forma atómica en SQLite y generar el comprobante .docx oficial en disco.
"""

from datetime import datetime, timedelta
import hashlib
import json
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
import urllib.error
import urllib.parse
import urllib.request

from docx import Document

from .mercado_libre_parser import _buscar_codigo
from .paths import application_root
from .stock.models import RegistroDescuento
from .stock_manager import normalizar_codigo

CONFIG_ML_PATH_DEFAULT = "data/config_mercadolibre.json"
API_BASE_URL = "https://api.mercadolibre.com"


def _obtener_ruta_config_ml(root_path: Path | None = None) -> Path:
    root = root_path or application_root()
    return root / CONFIG_ML_PATH_DEFAULT


def cargar_config_mercadolibre(root_path: Path | None = None) -> dict[str, Any]:
    """Carga la configuración de la API de Mercado Libre desde disco o variables de entorno."""
    ruta = _obtener_ruta_config_ml(root_path)
    res = {
        "app_id": "",
        "client_secret": "",
        "seller_id": "",
        "access_token": "",
        "refresh_token": "",
        "activo": False,
        "filtro_estado": "paid",
        "horas_atras": 48,
        "auto_descontar": False,
        "ultima_sincronizacion": "",
    }
    if ruta.exists():
        try:
            with open(ruta, "r", encoding="utf-8") as f:
                data = json.load(f)
                res = {
                    "app_id": str(data.get("app_id", "")).strip(),
                    "client_secret": str(data.get("client_secret", "")).strip(),
                    "seller_id": str(data.get("seller_id", "")).strip(),
                    "access_token": str(data.get("access_token", "")).strip(),
                    "refresh_token": str(data.get("refresh_token", "")).strip(),
                    "activo": bool(data.get("activo", False)),
                    "filtro_estado": str(data.get("filtro_estado", "paid")).strip() or "paid",
                    "horas_atras": int(data.get("horas_atras", 48)),
                    "auto_descontar": bool(data.get("auto_descontar", False)),
                    "ultima_sincronizacion": str(data.get("ultima_sincronizacion", "")),
                }
        except Exception:
            pass

    for clave, env_nombres in [
        ("app_id", ["STOCKAGENT_MELI_APP_ID", "MELI_APP_ID"]),
        ("client_secret", ["STOCKAGENT_MELI_CLIENT_SECRET", "MELI_CLIENT_SECRET"]),
        ("seller_id", ["STOCKAGENT_MELI_SELLER_ID", "MELI_SELLER_ID"]),
        ("access_token", ["STOCKAGENT_MELI_ACCESS_TOKEN", "MELI_ACCESS_TOKEN"]),
        ("refresh_token", ["STOCKAGENT_MELI_REFRESH_TOKEN", "MELI_REFRESH_TOKEN"]),
    ]:
        if not res.get(clave):
            for env_nombre in env_nombres:
                val = os.environ.get(env_nombre)
                if val:
                    res[clave] = val.strip()
                    break

    return res


def guardar_config_mercadolibre(config: dict[str, Any], root_path: Path | None = None) -> None:
    """Persiste la configuración de la API de Mercado Libre en disco."""
    ruta = _obtener_ruta_config_ml(root_path)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(config, f, indent=2, ensure_ascii=False)


def renovar_token_ml(
    config: dict[str, Any], root_path: Path | None = None
) -> tuple[bool, str, str]:
    """Renueva el Access Token de Mercado Libre utilizando el Refresh Token y Client Secret."""
    app_id = config.get("app_id", "").strip()
    client_secret = config.get("client_secret", "").strip()
    refresh_token = config.get("refresh_token", "").strip()

    if not app_id or not client_secret or not refresh_token:
        return False, "Faltan App ID, Client Secret o Refresh Token para renovar el token.", ""

    url = f"{API_BASE_URL}/oauth/token"
    payload = urllib.parse.urlencode({
        "grant_type": "refresh_token",
        "client_id": app_id,
        "client_secret": client_secret,
        "refresh_token": refresh_token,
    }).encode("utf-8")

    req = urllib.request.Request(
        url,
        data=payload,
        headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "Accept": "application/json",
            "User-Agent": "StockAgent/2.0",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            nuevo_access = str(data.get("access_token", ""))
            nuevo_refresh = str(data.get("refresh_token", refresh_token))
            user_id = str(data.get("user_id", ""))

            if not nuevo_access:
                return False, "Mercado Libre no retornó un access_token válido.", ""

            config["access_token"] = nuevo_access
            config["refresh_token"] = nuevo_refresh
            if user_id and not config.get("seller_id"):
                config["seller_id"] = user_id

            guardar_config_mercadolibre(config, root_path=root_path)
            return True, "Token renovado exitosamente.", nuevo_access

    except urllib.error.HTTPError as he:
        detalle = he.read().decode("utf-8", errors="ignore")
        return False, f"Error HTTP {he.code} al renovar token: {detalle}", ""
    except Exception as e:
        return False, f"Error de conexión al renovar token: {e}", ""


def verificar_credenciales_ml(
    config: dict[str, Any] | None = None, root_path: Path | None = None
) -> tuple[bool, str, dict[str, Any]]:
    """Comprueba la validez del Access Token consultando el perfil del usuario autenticado (/users/me)."""
    cfg = config or cargar_config_mercadolibre(root_path)
    token = cfg.get("access_token", "").strip()

    if not token:
        return False, "Falta ingresar el Access Token de Mercado Libre.", {}

    url = f"{API_BASE_URL}/users/me"
    req = urllib.request.Request(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "User-Agent": "StockAgent/2.0",
        },
        method="GET",
    )

    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            nickname = data.get("nickname", "Desconocido")
            user_id = data.get("id", "")
            email = data.get("email", "")

            # Guardar automáticamente el seller_id si no estaba cargado
            if user_id and not cfg.get("seller_id"):
                cfg["seller_id"] = str(user_id)
                guardar_config_mercadolibre(cfg, root_path=root_path)

            return True, f"Conexión exitosa con la cuenta '{nickname}' (ID: {user_id})", data

    except urllib.error.HTTPError as he:
        if he.code == 401 and cfg.get("refresh_token"):
            # Intentar renovación automática
            exito_renov, msg_renov, nuevo_tok = renovar_token_ml(cfg, root_path=root_path)
            if exito_renov:
                cfg["access_token"] = nuevo_tok
                return verificar_credenciales_ml(cfg, root_path=root_path)
            return False, f"Token expirado y falló la renovación: {msg_renov}", {}
        return False, f"Error HTTP {he.code} de Mercado Libre: {he.reason}", {}
    except Exception as e:
        return False, f"Error de red al conectar con Mercado Libre: {e}", {}


def consultar_ordenes_ml(
    config: dict[str, Any] | None = None,
    horas: int = 48,
    limit: int = 50,
    root_path: Path | None = None,
) -> tuple[bool, str, list[dict[str, Any]]]:
    """Descarga de la API de Mercado Libre las órdenes recientes pagadas."""
    cfg = config or cargar_config_mercadolibre(root_path)
    token = cfg.get("access_token", "").strip()
    if not token:
        return False, "No hay Access Token configurado.", []

    seller_id = cfg.get("seller_id", "").strip()
    if not seller_id:
        # Intentar obtener el seller_id de la cuenta
        ok_conn, _, user_data = verificar_credenciales_ml(cfg, root_path=root_path)
        if ok_conn and user_data.get("id"):
            seller_id = str(user_data["id"])
            cfg["seller_id"] = seller_id
        else:
            seller_id = "me"

    fecha_desde = datetime.now() - timedelta(hours=horas)
    fecha_desde_iso = fecha_desde.strftime("%Y-%m-%dT%H:%M:%S.000-03:00")

    params = {
        "seller": seller_id,
        "order.status": cfg.get("filtro_estado", "paid"),
        "order.date_created.from": fecha_desde_iso,
        "sort": "date_desc",
        "limit": str(min(limit, 50)),
    }

    url = f"{API_BASE_URL}/orders/search?{urllib.parse.urlencode(params)}"
    req = urllib.request.Request(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/json",
            "User-Agent": "StockAgent/2.0",
        },
        method="GET",
    )

    try:
        with urllib.request.urlopen(req, timeout=12) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            results = data.get("results", [])

            cfg["ultima_sincronizacion"] = datetime.now().isoformat()
            guardar_config_mercadolibre(cfg, root_path=root_path)

            return True, f"Se sincronizaron {len(results)} órdenes de Mercado Libre.", results

    except urllib.error.HTTPError as he:
        if he.code == 401 and cfg.get("refresh_token"):
            ok_renov, _, _ = renovar_token_ml(cfg, root_path=root_path)
            if ok_renov:
                return consultar_ordenes_ml(cfg, horas=horas, limit=limit, root_path=root_path)
        return False, f"Error HTTP {he.code} al consultar órdenes: {he.reason}", []
    except Exception as e:
        return False, f"Error al consultar órdenes de Mercado Libre: {e}", []


def mapear_orden_ml_a_items(
    orden: dict[str, Any],
    catalogo: dict[str, tuple[int, str, int]],
) -> dict[str, Any]:
    """Empareja los ítems de una orden de Mercado Libre con el catálogo de StockAgent.
    
    Estrategia de matching:
    1. SKU exacto del vendedor (seller_sku o seller_custom_field) contra el código de catálogo.
    2. Si no hay SKU o no coincide, fuzzy matching semántico del título contra los nombres del catálogo.
    """
    orden_id = str(orden.get("id", ""))
    ref_pedido = f"ML-{orden_id}"
    buyer = orden.get("buyer") or {}
    nombre_comprador = (
        f"{buyer.get('first_name', '')} {buyer.get('last_name', '')}".strip()
        or str(buyer.get("nickname", "Comprador Mercado Libre"))
    )
    fecha_creacion = str(orden.get("date_created", ""))[:19].replace("T", " ")
    shipping = orden.get("shipping") or {}
    shipping_id = str(shipping.get("id", ""))

    descripciones_catalogo: dict[str, list[str]] = {}
    for cod, data in catalogo.items():
        nom = data[1]
        if nom:
            descripciones_catalogo.setdefault(cod, []).append(nom)

    items_coincidentes: list[dict[str, Any]] = []
    items_no_coincidentes: list[dict[str, Any]] = []

    for item_elem in orden.get("order_items", []):
        it_obj = item_elem.get("item") or {}
        sku_candidato = normalizar_codigo(
            str(it_obj.get("seller_sku") or it_obj.get("seller_custom_field") or "").strip()
        )
        titulo = str(it_obj.get("title", "")).strip()
        cantidad = int(item_elem.get("quantity", 1))
        precio_unitario = float(item_elem.get("unit_price", 0.0))

        codigo_encontrado: str | None = None
        metodo = ""

        # 1. Matching por SKU
        if sku_candidato and sku_candidato in catalogo:
            codigo_encontrado = sku_candidato
            metodo = "SKU"
        else:
            # 2. Matching fuzzy por título
            codigo_fuzzy, _ = _buscar_codigo(titulo, descripciones_catalogo)
            if codigo_fuzzy and codigo_fuzzy in catalogo:
                codigo_encontrado = codigo_fuzzy
                metodo = "Fuzzy Título"

        if codigo_encontrado:
            fila, nombre_prod, stock_act = catalogo[codigo_encontrado]
            items_coincidentes.append({
                "codigo": codigo_encontrado,
                "nombre": nombre_prod,
                "cantidad": cantidad,
                "stock_actual": stock_act,
                "precio": precio_unitario,
                "fila": fila,
                "metodo": metodo,
                "stock_suficiente": stock_act >= cantidad,
            })
        else:
            items_no_coincidentes.append({
                "titulo": titulo,
                "sku": sku_candidato,
                "cantidad": cantidad,
                "precio": precio_unitario,
            })

    total_unidades = sum(it["cantidad"] for it in items_coincidentes)
    total_importe = sum(it["cantidad"] * it["precio"] for it in items_coincidentes)
    todos_coinciden = len(items_no_coincidentes) == 0 and len(items_coincidentes) > 0
    stock_suficiente = all(it["stock_suficiente"] for it in items_coincidentes)

    return {
        "orden_id": orden_id,
        "ref_pedido": ref_pedido,
        "cliente": nombre_comprador,
        "fecha": fecha_creacion,
        "shipping_id": shipping_id,
        "items_coincidentes": items_coincidentes,
        "items_no_coincidentes": items_no_coincidentes,
        "total_unidades": total_unidades,
        "total_importe": total_importe,
        "todos_coinciden": todos_coinciden,
        "stock_suficiente": stock_suficiente,
    }


def generar_documento_orden_ml_docx(orden_procesada: dict[str, Any], ruta_salida: Path) -> Path:
    """Genera un archivo .docx oficial para resguardar la orden de Mercado Libre en pedidos_archivos."""
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(f"Mercado Libre · Orden {orden_procesada['orden_id']}", level=1)

    p_meta = doc.add_paragraph()
    p_meta.add_run("Cliente / Comprador: ").bold = True
    p_meta.add_run(f"{orden_procesada.get('cliente', 'Mercado Libre')}\n")
    p_meta.add_run("Fecha: ").bold = True
    p_meta.add_run(f"{orden_procesada.get('fecha', '')}\n")
    p_meta.add_run("Identificador de Envío: ").bold = True
    p_meta.add_run(f"{orden_procesada.get('shipping_id', 'N/A')}\n")
    p_meta.add_run("Total de Artículos: ").bold = True
    p_meta.add_run(f"{orden_procesada.get('total_unidades', 0)} unidades\n")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Código (SKU)"
    hdr_cells[1].text = "Descripción de Producto"
    hdr_cells[2].text = "Cant."
    hdr_cells[3].text = "Precio Unitario"

    for it in orden_procesada.get("items_coincidentes", []):
        row_cells = table.add_row().cells
        row_cells[0].text = str(it.get("codigo", ""))
        row_cells[1].text = str(it.get("nombre", ""))
        row_cells[2].text = str(it.get("cantidad", 0))
        row_cells[3].text = f"${it.get('precio', 0.0):,.2f}"

    doc.save(str(ruta_salida))
    return ruta_salida


def descontar_orden_ml_api(
    orden_procesada: dict[str, Any],
    stock_manager: Any,
    order_store: Any,
    armador: str = "Operario",
    root_path: Path | None = None,
) -> tuple[bool, str]:
    """Aplica el descuento de stock atómico de una orden de Mercado Libre y la registra en OrderStore."""
    ref_pedido = orden_procesada["ref_pedido"]
    orden_id = orden_procesada["orden_id"]

    # Validar si ya fue procesada
    if order_store.existe_nota_pedido(ref_pedido):
        return False, f"La orden {ref_pedido} ya fue descontada previamente en el sistema."

    hash_orden = hashlib.sha256(f"ML_API_ORDER_{orden_id}".encode()).hexdigest()
    dup = order_store.buscar_pedido_por_hash(hash_orden)
    if dup:
        return False, f"La orden {ref_pedido} ya fue procesada anteriormente ({dup.created_at[:16]})."

    items = orden_procesada.get("items_coincidentes", [])
    if not items:
        return False, f"La orden {ref_pedido} no tiene ítems válidos para descontar."

    # Generar el archivo .docx de respaldo en data/pedidos_archivos/
    root = root_path or application_root()
    carpeta_pedidos = root / "data" / "pedidos_archivos"
    nombre_doc = f"ML_API_{orden_id}.docx"
    ruta_doc = carpeta_pedidos / nombre_doc
    generar_documento_orden_ml_docx(orden_procesada, ruta_doc)

    # Preparar descuentos para StockManager
    descuentos: list[RegistroDescuento] = []
    for it in items:
        cod = it["codigo"]
        nom = it["nombre"]
        cant = it["cantidad"]
        stk_ant = it["stock_actual"]
        stk_nuevo = int(stk_ant - cant)
        fila = it.get("fila", 0)

        descuentos.append(
            RegistroDescuento(
                codigo=cod,
                nombre_producto=nom,
                fila_excel=fila,
                stock_anterior=stk_ant,
                cantidad_descontada=cant,
                stock_nuevo=stk_nuevo,
                es_nuevo_agregado=False,
            )
        )

    # Actualizar stock base atómico
    stock_manager.actualizar_stock_base(
        descuentos=descuentos,
        sincronizar_excel=True,
        nombre_archivo=nombre_doc,
        usuario=armador,
    )

    total_desc = orden_procesada["total_unidades"]
    resumen_items = f"{len(items)} productos ({total_desc} un.)"
    items_descontados = [(it["codigo"], it["nombre"], it["cantidad"]) for it in items]

    # Registrar en OrderStore
    order_store.registrar_pedido(
        nota_pedido=ref_pedido,
        numero_pedido=ref_pedido,
        cliente=orden_procesada.get("cliente", "Mercado Libre"),
        bultos=1,
        logistica="Transporte",
        armador=armador,
        tipo_archivo="docx",
        nombre_archivo=nombre_doc,
        agregado=False,
        articulos_descontados=total_desc,
        resumen_items=resumen_items,
        archivo_descuento="",
        items=items_descontados,
        archivo_hash=hash_orden,
    )

    return True, f"Orden {ref_pedido} descontada con éxito ({total_desc} unidades)."
