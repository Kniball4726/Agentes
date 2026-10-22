"""Lee pedidos de Mercado Libre en DOCX y asocia productos con el catálogo."""

import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document

from .remito_parser import ItemRemito, ResultadoRemito

_UMBRAL_SIMILITUD = 0.84
_MARGEN_MINIMO_SIMILITUD = 0.08
_SABORES = {
    "banana", "coco", "chocolate", "dulce", "frambuesa", "frutilla",
    "leche", "limon", "naranja", "vainilla",
}


def _normalizar_descripcion(descripcion: str) -> str:
    texto = unicodedata.normalize("NFKD", descripcion.casefold())
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return " ".join(re.sub(r"[^\w]+", " ", texto).split())


def _normalizar_para_busqueda(descripcion: str) -> str:
    texto = _normalizar_descripcion(descripcion)
    equivalencias = (
        (r"\bultra\s+tech\b", "ut"),
        (r"\bvitatech\b", "vt"),
        (r"\bcreatina\b", "creatine"),
        (r"\bcapsulas?\b", "caps"),
        (r"\bcomprimidos?\b", "comp"),
        (r"\bgramos?\b|\bgrs?\b", "g"),
        (r"\bunidades\b", "unidad"),
    )
    for patron, reemplazo in equivalencias:
        texto = re.sub(patron, reemplazo, texto)
    texto = re.sub(r"\b(de|del|x)\b", " ", texto)
    return " ".join(texto.split())


def _buscar_codigo(
    descripcion: str,
    descripciones_catalogo: Dict[str, List[str]]
) -> Tuple[str | None, str]:
    clave = _normalizar_para_busqueda(descripcion)
    catalogo_normalizado = {
        codigo: [_normalizar_para_busqueda(nombre) for nombre in nombres]
        for codigo, nombres in descripciones_catalogo.items()
    }
    codigos = [
        codigo for codigo, nombres in catalogo_normalizado.items()
        if clave in nombres
    ]
    if len(codigos) == 1:
        return codigos[0], ""
    if codigos:
        return None, "Descripción coincidente con más de un código del catálogo"

    tokens = clave.split()
    if not tokens:
        return None, "No se encontró una coincidencia exacta en el catálogo"

    tokens_producto = set(tokens) - {"ut", "vt"}
    if tokens_producto:
        codigos_por_terminos = [
            codigo
            for codigo, nombres in catalogo_normalizado.items()
            if any(
                tokens_producto.issubset(set(nombre.split()) - {"ut", "vt"})
                for nombre in nombres
            )
        ]
        if len(codigos_por_terminos) == 1:
            return codigos_por_terminos[0], ""
        if codigos_por_terminos:
            return None, "Coincidencia aproximada ambigua entre varios productos"

    numeros_consulta = {token for token in tokens if token.isdigit()}
    sabores_consulta = set(tokens) & _SABORES
    puntuaciones = {}
    for codigo, nombres in catalogo_normalizado.items():
        puntuacion = 0.0
        for nombre in nombres:
            tokens_nombre = nombre.split()
            numeros_producto = {token for token in tokens_nombre if token.isdigit()}
            sabores_producto = set(tokens_nombre) & _SABORES
            if (
                not tokens_nombre
                or tokens_nombre[0] != tokens[0]
                or not numeros_producto.issubset(numeros_consulta)
                or (
                    sabores_consulta
                    and sabores_producto
                    and sabores_consulta.isdisjoint(sabores_producto)
                )
            ):
                continue
            puntuacion = max(
                puntuacion,
                SequenceMatcher(None, clave, nombre).ratio(),
            )
        if puntuacion >= _UMBRAL_SIMILITUD:
            puntuaciones[codigo] = puntuacion

    mejores = sorted(puntuaciones.items(), key=lambda item: item[1], reverse=True)
    if mejores:
        if (
            len(mejores) > 1
            and mejores[0][1] - mejores[1][1] < _MARGEN_MINIMO_SIMILITUD
        ):
            return None, "Coincidencia aproximada ambigua entre varios productos"
        return mejores[0][0], ""
    return None, "No se encontró una coincidencia suficientemente clara en el catálogo"


def _extraer_producto_numerado(linea: str) -> Tuple[int, str] | None:
    coincidencia = re.match(r"^\s*(\d+)[.)]?\s+(.+?)\s*$", linea)
    if not coincidencia:
        return None
    return int(coincidencia.group(1)), coincidencia.group(2)


def parsear_pedido_mercado_libre(
    ruta_docx: str | Path,
    productos_catalogo: List[Tuple[str, str]]
) -> ResultadoRemito:
    """Procesa líneas de producto numeradas y usa su número inicial como cantidad."""
    path = Path(ruta_docx)
    documento = Document(str(path))
    descripciones_catalogo: Dict[str, List[str]] = {}
    for codigo, nombre in productos_catalogo:
        if nombre:
            nombres = descripciones_catalogo.setdefault(codigo, [])
            if nombre not in nombres:
                nombres.append(nombre)

    registros: List[Tuple[int, str]] = []
    for parrafo in documento.paragraphs:
        for linea in parrafo.text.splitlines():
            producto = _extraer_producto_numerado(linea)
            if producto:
                registros.append(producto)

    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [
                [linea.strip() for linea in celda.text.splitlines() if linea.strip()]
                for celda in fila.cells
            ]
            productos_fila: List[Tuple[int, str]] = []
            indice = 0
            while indice < len(celdas):
                for linea in celdas[indice]:
                    producto = _extraer_producto_numerado(linea)
                    if producto:
                        productos_fila.append(producto)
                    elif re.fullmatch(r"\s*\d+\s*[.)]?\s*", linea):
                        siguiente = indice + 1
                        while siguiente < len(celdas) and not celdas[siguiente]:
                            siguiente += 1
                        if siguiente < len(celdas):
                            descripcion = " ".join(celdas[siguiente])
                            cantidad = int(re.search(r"\d+", linea).group())
                            productos_fila.append((cantidad, descripcion))
                indice += 1
            registros.extend(dict.fromkeys(productos_fila))

    items: List[ItemRemito] = []
    sin_coincidencia: Dict[str, Tuple[int, str, str]] = {}
    for cantidad, descripcion in registros:
        codigo, motivo = _buscar_codigo(descripcion, descripciones_catalogo)
        if codigo is None:
            clave = _normalizar_descripcion(descripcion)
            cantidad_acumulada, descripcion_original, motivo_original = (
                sin_coincidencia.get(clave, (0, descripcion, motivo))
            )
            sin_coincidencia[clave] = (
                cantidad_acumulada + cantidad,
                descripcion_original,
                motivo_original,
            )
            continue

        items.append(ItemRemito(
            codigo=codigo,
            descripcion=descripcion,
            cantidad=cantidad,
            origen_remito=f"ML-{path.stem}"
        ))

    total_unidades = sum(cantidad for cantidad, _ in registros)
    return ResultadoRemito(
        archivo=path.name,
        numero_remito=f"ML-{path.stem}",
        fecha="Mercado Libre",
        cliente="Mercado Libre",
        articulos_declarados=total_unidades,
        articulos_extraidos=sum(item.cantidad for item in items),
        verificacion_ok=not sin_coincidencia,
        items=items,
        descripciones_sin_coincidencia=[
            (f"{cantidad} {descripcion}", motivo)
            for cantidad, descripcion, motivo in sin_coincidencia.values()
        ]
    )
