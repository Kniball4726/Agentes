"""
Módulo para parsear planillas Excel (.xlsx, .xls) de relevamiento de Stock Físico.
Extrae los códigos de los productos y las cantidades físicas contadas por el operario (incluyendo 0).
"""

from pathlib import Path
from typing import Dict, List, Tuple
import openpyxl
import xlrd
from .remito_parser import normalizar_codigo


HEADERS_CODIGO = ["CODIGO", "CÓDIGO", "COD", "SKU", "ARTICULO", "ARTÍCULO", "ID"]
HEADERS_DESCRIPCION = ["PRODUCTOS", "PRODUCTO", "DESCRIPCION", "DESCRIPCIÓN", "NOMBRE", "DETALLE", "TITULO", "TÍTULO"]
HEADERS_FISICO_PRIORIDAD = ["STOCK FISICO", "STOCK FÍSICO", "FISICO", "FÍSICO", "CONTEO", "REAL", "RELEVADO"]
HEADERS_CANTIDAD_FALLBACK = ["CANTIDAD", "CANT", "UNIDADES", "STOCK", "EXISTENCIA"]


def _detectar_columnas(filas_texto: List[List[str]]) -> Tuple[int, int, int, int]:
    """
    Detecta índices (base 0) de (col_cod, col_desc, col_cant, fila_inicio_datos).
    Prioriza columnas con nombre explícito de stock físico/conteo sobre stock general.
    """
    for r_idx, fila_vals in enumerate(filas_texto):
        col_cod = -1
        col_desc = -1
        col_cant_prio = -1
        col_cant_fallback = -1

        for c_idx, val in enumerate(fila_vals):
            if not val:
                continue
            if any(k in val for k in HEADERS_CODIGO) and col_cod == -1:
                col_cod = c_idx
            elif any(k in val for k in HEADERS_DESCRIPCION) and col_desc == -1:
                col_desc = c_idx
            elif any(k in val for k in HEADERS_FISICO_PRIORIDAD) and col_cant_prio == -1:
                col_cant_prio = c_idx
            elif any(k in val for k in HEADERS_CANTIDAD_FALLBACK):
                col_cant_fallback = c_idx

        col_cant = col_cant_prio if col_cant_prio != -1 else col_cant_fallback
        if col_cod != -1 and col_cant != -1 and col_cod != col_cant:
            return col_cod, col_desc, col_cant, r_idx + 1

    return -1, -1, -1, 0


def parsear_excel_stock_fisico(
    ruta_archivo: str | Path,
) -> Tuple[Dict[str, int], List[Tuple[str, str, int]]]:
    """
    Parsea una planilla Excel (.xlsx o .xls) de inventario físico.

    Retorna:
        - items_consolidados: Dict[codigo, stock_fisico_total] (admite 0 o mayores).
        - detalle: List[(codigo, descripcion, stock_fisico)] fila por fila.
    """
    path = Path(ruta_archivo)
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el archivo de stock físico: {path}")

    ext = path.suffix.lower()
    if ext not in (".xlsx", ".xls"):
        raise ValueError(
            f"Formato no soportado para stock físico: {ext}. Debe ser .xlsx o .xls"
        )

    items_crudos: List[Tuple[str, str, int]] = []

    if ext == ".xls":
        libro = xlrd.open_workbook(str(path))
        hoja = libro.sheet_by_index(0)

        primeras_filas = [
            [str(hoja.cell_value(r, c)).strip().upper() for c in range(hoja.ncols)]
            for r in range(min(20, hoja.nrows))
        ]
        col_cod, col_desc, col_cant, fila_inicio = _detectar_columnas(primeras_filas)

        if col_cod == -1 or col_cant == -1:
            col_cod = 0
            col_desc = 1 if hoja.ncols > 2 else -1
            col_cant = 2 if hoja.ncols > 2 else (1 if hoja.ncols > 1 else 0)
            fila_inicio = 0

        for r in range(fila_inicio, hoja.nrows):
            val_cod = hoja.cell_value(r, col_cod)
            val_cant = hoja.cell_value(r, col_cant)
            val_desc = hoja.cell_value(r, col_desc) if col_desc != -1 else ""

            if val_cod is None or str(val_cod).strip() == "":
                continue

            str_cod = str(val_cod).strip()
            if str_cod.upper() in ("TOTAL", "TOTALES", "SUBTOTAL", "RESUMEN", "CODIGO", "CÓDIGO"):
                continue

            if val_cant is None or str(val_cant).strip() == "":
                continue

            cod = normalizar_codigo(str_cod)
            try:
                cant = int(float(val_cant))
            except (ValueError, TypeError):
                continue

            if cant >= 0:
                items_crudos.append((cod, str(val_desc).strip(), cant))

    else:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        try:
            hoja = wb.active
            if hoja is None:
                return {}, []

            primeras_filas = [
                [
                    str(hoja.cell(r, c).value or "").strip().upper()
                    for c in range(1, hoja.max_column + 1)
                ]
                for r in range(1, min(20, hoja.max_row + 1))
            ]
            col_cod_0, col_desc_0, col_cant_0, fila_inicio_0 = _detectar_columnas(
                primeras_filas
            )

            if col_cod_0 == -1 or col_cant_0 == -1:
                col_cod = 1
                col_desc = 2 if hoja.max_column > 2 else -1
                col_cant = 3 if hoja.max_column > 2 else (2 if hoja.max_column > 1 else 1)
                fila_inicio = 1
            else:
                col_cod = col_cod_0 + 1
                col_desc = (col_desc_0 + 1) if col_desc_0 != -1 else -1
                col_cant = col_cant_0 + 1
                fila_inicio = fila_inicio_0 + 1

            for r in range(fila_inicio, hoja.max_row + 1):
                val_cod = hoja.cell(r, col_cod).value
                val_cant = hoja.cell(r, col_cant).value
                val_desc = hoja.cell(r, col_desc).value if col_desc != -1 else ""

                if val_cod is None or str(val_cod).strip() == "":
                    continue

                str_cod = str(val_cod).strip()
                if str_cod.upper() in ("TOTAL", "TOTALES", "SUBTOTAL", "RESUMEN", "CODIGO", "CÓDIGO"):
                    continue

                if val_cant is None or str(val_cant).strip() == "":
                    continue

                cod = normalizar_codigo(str_cod)
                try:
                    cant = int(float(val_cant))
                except (ValueError, TypeError):
                    continue

                if cant >= 0:
                    items_crudos.append((cod, str(val_desc or "").strip(), cant))
        finally:
            wb.close()

    items_consolidados: Dict[str, int] = {}
    for cod, _, cant in items_crudos:
        items_consolidados[cod] = items_consolidados.get(cod, 0) + cant

    return items_consolidados, items_crudos
